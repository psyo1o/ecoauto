# -*- coding: utf-8 -*-
"""
측정인.kr 자동 비교 시스템 - FINAL + 파일명 자동 생성 + PDF 다운로드/하이퍼링크 + NG 빨간색 표시
"""
import warnings
warnings.filterwarnings("ignore")  # 무조건 모든 경고 차단 (조건 없음)
warnings.showwarning = lambda *args, **kwargs: None
import os
import re
import time
from datetime import datetime

from selenium import webdriver
from selenium.webdriver.common.by import By
from selenium.webdriver.common.action_chains import ActionChains
from selenium.webdriver.support.ui import WebDriverWait
from selenium.webdriver.support import expected_conditions as EC
from selenium.common.exceptions import TimeoutException

from openpyxl import Workbook, load_workbook
from openpyxl.utils import range_boundaries
from openpyxl.styles import PatternFill
from openpyxl.formatting.rule import FormulaRule

# ============================================================
# 공통 유틸 모듈 (모듈화)
# ============================================================
from selenium_utils import safe_click, wait_el as wait_until_exists, close_popup, set_date_js, wait
from format_utils import (
    format_time as trim_time_to_hm,
    to_float1,
    to_float2,
    parse_datetime_text,
    normalize_tab1_select_field,
    facility_labels_match,
)
from data_utils import norm_ymd, sample_to_datestr, clean_leading_mark
from excel_utils import find_sheet_by_candidates as _find_sheet_by_candidates_openpyxl
from file_utils import find_best_matching_file as _find_best_file_util, is_fugitive_dust_file
from measin_utils import (
    login, search_date, wait_grid_loaded, get_samples_current_page,
    go_back_to_list, collect_samples_from_files,
    open_detail_with_session_recovery, recover_site_session, is_field_list_ready,
    ensure_detail_page_for_tab1, reopen_sample_from_search,
    ensure_logged_in_or_recover, is_logged_out,
    verify_tab4_list_status,
    MAX_SAMPLE_DETAIL_RETRY,
    LOGIN_URL, FIELD_URL, NAS_BASE, NAS_DIRS
)
from excel_utils import find_sheet_by_candidates, parse_measuring_record, autofit_columns
from realgrid_utils import rg_api_read_data
from log_utils import log_error
from cancel_utils import is_cancelled
from config import MEASIN_REVIEW, MEASIN_PDF_DIR, MEASIN_PHOTO_DIR
from measin_constants import (
    SKIP_VOL_AND_SPEED, SKIP_SPEED_ONLY, DUST_SKIP_FIELDS,
    SM3_ITEMS as sm3_items,
    SEL_DATE, SEL_START_TIME, SEL_END_TIME, SEL_WEATHER, SEL_EMIS_FAC,
    SEL_O2_STD, SEL_O2_MEAS, SEL_GAS_VOL_PRE, SEL_GAS_VOL_POST,
    SEL_MOISTURE, SEL_GAS_TEMP, SEL_GAS_SPEED
)

# ------------------------------------------------------------
# 설정
# ------------------------------------------------------------

# 구버전 호환(루트). 실제 저장은 sample별 연/월 폴더 사용.
PDF_DIR = MEASIN_PDF_DIR
PHOTO_DIR = MEASIN_PHOTO_DIR
if not os.path.isdir(PDF_DIR):
    os.makedirs(PDF_DIR)
if not os.path.isdir(PHOTO_DIR):
    os.makedirs(PHOTO_DIR)

PDF_MAP = {}
PHOTO_MAP = {}  # sample_no -> [{"idx", "path", "shot_at"}, ...]
COMPANY_MAP = {}   # ★ 추가: 시료번호 -> 업소명(표시용)

# 목록 RealGrid '상태' 열 기대값 (탭1·2·3 자료수집 후 목록 복귀 시 확인)
CHECK_LIST_EXPECTED_STATUS = "측정분석결과 입력중"


def _ym_folder_from_sample(sample_no: str) -> tuple[str, str]:
    """시료번호 → (연도, 'N월'). 예: A2601164 → ('2026', '1월')."""
    ds = sample_to_datestr(str(sample_no or "").strip())
    if ds:
        try:
            yyyy, mm, _dd = ds.split("-")
            return yyyy, f"{int(mm)}월"
        except Exception:
            pass
    now = datetime.now()
    return str(now.year), f"{now.month}월"


def media_dir_for_sample(kind: str, sample_no: str) -> str:
    """
    PDF/현장사진 저장 폴더.
    예: ...\\3.측정인 검토\\PDF\\2026\\9월
         ...\\3.측정인 검토\\현장사진\\2026\\9월
    """
    yyyy, mlabel = _ym_folder_from_sample(sample_no)
    base = PDF_DIR if kind == "PDF" else PHOTO_DIR
    path = os.path.join(base, yyyy, mlabel)
    os.makedirs(path, exist_ok=True)
    return path


# ============================================================
# RealGrid 비교 예외 → measin_constants.py 에서 import 완료
# ============================================================

# ------------------------------------------------------------
# 공통 유틸 (wait → selenium_utils, clean_leading_mark → data_utils)
# ------------------------------------------------------------


def init_driver():
    """eco_check 전용 드라이버 초기화 (PDF 다운로드 경로 설정 포함)"""
    from selenium_utils import init_driver as _base_init
    d = _base_init()
    # 초기 다운로드 경로(시료별로는 download_pdf에서 재설정)
    try:
        d.execute_cdp_cmd(
            "Page.setDownloadBehavior",
            {"behavior": "allow", "downloadPath": PDF_DIR}
        )
    except:
        pass
    return d


def _wait_new_pdf(download_dir, before_set, timeout=60):
    """
    클릭 직전(before_set) 대비 새로 생긴 PDF를 찾아서
    .crdownload가 사라지고 완성된 파일만 반환
    """
    t0 = time.time()

    while time.time() - t0 < timeout:
        now = set(os.listdir(download_dir))

        # 새로 생긴 파일 후보
        added = list(now - before_set)

        # 크롬 다운로드 진행중이면 .crdownload 존재
        crs = [f for f in added if f.lower().endswith(".crdownload")]
        pdfs = [f for f in added if f.lower().endswith(".pdf")]

        if crs:
            time.sleep(0.3)
            continue

        if pdfs:
            paths = [os.path.join(download_dir, f) for f in pdfs]
            latest = max(paths, key=os.path.getmtime)
            # 파일 쓰기 마무리 안정화
            time.sleep(0.5)
            return latest

        time.sleep(0.3)

    return ""


# ------------------------------------------------------------
# PDF 다운로드
# ------------------------------------------------------------
def download_pdf(driver, sample_no):
    """
    PDF 다운로드 버튼 클릭 후
    '이번 클릭으로 새로 생성된 PDF'만 잡아서
    ...\\3.측정인 검토\\PDF\\{연도}\\{월}\\{sample_no}.pdf 로 저장
    """
    print(f"   [PDF] 다운로드 시도: {sample_no}")

    out_dir = media_dir_for_sample("PDF", sample_no)
    target_path = os.path.join(out_dir, f"{sample_no}.pdf")
    pdf_btn_sel = "#fileArea > section > div > div.row.fr > input:nth-child(3)"

    if os.path.isfile(target_path):
        try:
            os.remove(target_path)
        except Exception:
            pass

    try:
        driver.execute_cdp_cmd(
            "Page.setDownloadBehavior",
            {"behavior": "allow", "downloadPath": out_dir},
        )
    except Exception as e:
        print(f"   ⚠ 다운로드 경로 설정 실패: {e}")

    try:
        before = set(os.listdir(out_dir))
    except Exception as e:
        print(f"   ❌ PDF_DIR 접근 실패: {e}")
        return ""

    if not safe_click(driver, pdf_btn_sel):
        print("   ❌ PDF 버튼 클릭 실패")
        return ""

    new_pdf = _wait_new_pdf(out_dir, before, timeout=60)
    if not new_pdf or not os.path.isfile(new_pdf):
        print("   ❌ PDF 다운로드 완료/파일 탐지 실패")
        return ""

    try:
        if os.path.abspath(new_pdf) != os.path.abspath(target_path):
            os.replace(new_pdf, target_path)
        print(f"   ✔ PDF 저장: {target_path}")
        return target_path
    except Exception as e:
        print(f"   ❌ PDF 이름 변경 실패: {e}")
        return new_pdf


# ------------------------------------------------------------
# 탭3 현장사진 저장 (#photo0~2)
# ------------------------------------------------------------
def _img_element_to_png_bytes(driver, img_el) -> bytes:
    """이미 로드된 <img>를 canvas로 PNG 바이트 추출."""
    try:
        data_url = driver.execute_script(
            """
            var img = arguments[0];
            if (!img) return '';
            if (!img.complete || !img.naturalWidth) return '';
            var c = document.createElement('canvas');
            c.width = img.naturalWidth;
            c.height = img.naturalHeight;
            var ctx = c.getContext('2d');
            ctx.drawImage(img, 0, 0);
            return c.toDataURL('image/png');
            """,
            img_el,
        )
        if not data_url or not str(data_url).startswith("data:image"):
            return b""
        import base64
        b64 = str(data_url).split(",", 1)[1]
        return base64.b64decode(b64)
    except Exception:
        return b""


def _download_url_with_driver_cookies(driver, url: str) -> bytes:
    """Selenium 쿠키로 상대/절대 URL 다운로드 (requests 없을 때 urllib)."""
    try:
        from urllib.parse import urljoin, urlparse
        from urllib.request import Request, build_opener, HTTPCookieProcessor
        from http.cookiejar import CookieJar

        if not url:
            return b""
        if url.startswith("/"):
            parsed = urlparse(driver.current_url)
            url = f"{parsed.scheme}://{parsed.netloc}{url}"

        jar = CookieJar()
        opener = build_opener(HTTPCookieProcessor(jar))
        # CookieJar에 selenium 쿠키 주입은 번거로워 header로 직접
        cookie_hdr = "; ".join(
            f"{c['name']}={c['value']}" for c in driver.get_cookies()
        )
        req = Request(url, headers={"Cookie": cookie_hdr, "User-Agent": "Mozilla/5.0"})
        with opener.open(req, timeout=30) as resp:
            return resp.read()
    except Exception:
        return b""


def download_field_photos(driver, sample_no: str, shot_times: list | None = None) -> list[dict]:
    """
    탭3 #photo0~#photo2 저장.
    경로: ...\\3.측정인 검토\\현장사진\\{연도}\\{월}\\{sample_no}_PIC1.png ...
    반환: [{"idx":1, "path":..., "shot_at":...}, ...]
    """
    print(f"   [현장사진] 저장 시도: {sample_no}")
    out_dir = media_dir_for_sample("현장사진", sample_no)

    times = list(shot_times or [])
    out: list[dict] = []
    for i in range(3):
        pic_no = i + 1
        shot_at = times[i] if i < len(times) else ""
        item = {"idx": pic_no, "path": "", "shot_at": shot_at, "src": ""}
        try:
            img = driver.find_element(By.CSS_SELECTOR, f"#photo{i}")
            try:
                driver.execute_script(
                    "arguments[0].scrollIntoView({block:'center'});", img
                )
            except Exception:
                pass
            for _ in range(15):
                ready = driver.execute_script(
                    "return !!(arguments[0].complete && arguments[0].naturalWidth > 0);",
                    img,
                )
                if ready:
                    break
                time.sleep(0.2)

            src = (img.get_attribute("src") or "").strip()
            item["src"] = src
            if not src:
                print(f"   ⚠ photo{i}: src 없음")
                out.append(item)
                continue

            raw = _img_element_to_png_bytes(driver, img)
            ext = ".png"
            if not raw:
                raw = _download_url_with_driver_cookies(driver, src)
                from urllib.parse import urlparse
                path_ext = os.path.splitext(urlparse(src).path)[1]
                if path_ext.lower() in (".png", ".jpg", ".jpeg", ".gif", ".webp"):
                    ext = path_ext.lower()
                    if ext == ".jpeg":
                        ext = ".jpg"

            if not raw:
                print(f"   ⚠ photo{i}: 다운로드 실패")
                out.append(item)
                continue

            target = os.path.join(out_dir, f"{sample_no}_PIC{pic_no}{ext}")
            if os.path.isfile(target):
                try:
                    os.remove(target)
                except Exception:
                    pass
            with open(target, "wb") as f:
                f.write(raw)
            item["path"] = target
            print(f"   ✔ 현장사진 저장: {target}" + (f" ({shot_at})" if shot_at else ""))
        except Exception as e:
            print(f"   ⚠ photo{i} 처리 실패: {e}")
        out.append(item)
    return out


# ------------------------------------------------------------
# 사이트 데이터 수집
# ------------------------------------------------------------
def gv(driver, selector):
    try:
        el = driver.find_element(By.CSS_SELECTOR, selector)
        v = el.get_attribute("value")
        if not v:
            v = el.text
        return (v or "").strip()
    except:
        return ""


def click_tab(driver, tab_id) -> bool:
    try:
        el = WebDriverWait(driver, 10).until(
            EC.element_to_be_clickable((By.CSS_SELECTOR, f"a#{tab_id}"))
        )
        driver.execute_script("arguments[0].scrollIntoView(true);", el)
        time.sleep(0.2)
        driver.execute_script("arguments[0].click();", el)

        if tab_id == "ui-id-1":
            WebDriverWait(driver, 10).until(
                EC.visibility_of_element_located((By.CSS_SELECTOR, "#machineDiv"))
            )
        elif tab_id == "ui-id-2":
            WebDriverWait(driver, 10).until(
                EC.visibility_of_element_located((By.CSS_SELECTOR, "#meas_start_time"))
            )
        elif tab_id == "ui-id-3":
            WebDriverWait(driver, 10).until(
                EC.visibility_of_element_located((By.CSS_SELECTOR, "td#officer_dt"))
            )
        time.sleep(0.3)
        return True
    except Exception:
        print(" ❌ 탭 전환 실패:", tab_id)
        return False


def get_weather_text(driver):
    """기상 select.meas_wthr 선택값 텍스트."""
    try:
        s = driver.find_element(By.CSS_SELECTOR, SEL_WEATHER)
        v = (s.get_attribute("value") or "").strip()
        if not v:
            return ""
        try:
            op = s.find_element(By.CSS_SELECTOR, f"option[value='{v}']")
            return (op.text or v).strip()
        except Exception:
            return v
    except Exception:
        # 구형 input 호환
        return gv(
            driver,
            "#idWHArea > div > div:nth-child(2) > fieldset > label.col.col-12 "
            "> table > tbody > tr > td:nth-child(1) > input",
        )


def get_wind_direction_text(driver):
    try:
        s = driver.find_element(
            By.CSS_SELECTOR,
            "#idWHArea > div > div:nth-child(2) > fieldset > label.col.col-12 "
            "> table > tbody > tr > td:nth-child(5) > select"
        )
        v = s.get_attribute("value")
        op = s.find_element(By.CSS_SELECTOR, f"option[value='{v}']")
        return op.text.strip()
    except:
        return ""


_DT_RE = re.compile(r"\d{4}-\d{2}-\d{2}\s+\d{2}:\d{2}")

def _extract_time(txt: str) -> str:
    """'촬영일시: 2026-01-09 08:35' 같은 텍스트에서 '2026-01-09 08:35'만 뽑기"""
    if not txt:
        return ""
    m = _DT_RE.search(txt)
    if m:
        return m.group(0)
    # fallback
    if ":" in txt:
        return txt.split(":", 1)[1].strip()
    return txt.strip()

def get_mobile_times(driver):
    result = {
        "환경기술인입력일시": "",
        "GPS위치확인일시": "",
        "촬영일시목록": []
    }

    # ------------------------------------------------------------
    # 1) 환경기술인 입력일시: common-work-info 영역의 "입력일시"만
    # ------------------------------------------------------------
    try:
        env_els = driver.find_elements(
            By.XPATH,
            "//div[contains(@class,'common-work-info')]"
            "//td[@id='officer_dt' and contains(normalize-space(.),'입력일시')]"
        )
        if env_els:
            result["환경기술인입력일시"] = _extract_time(env_els[0].text.strip())
    except:
        pass

    # ------------------------------------------------------------
    # 2) 촬영일시 목록: pic_area 영역의 "촬영일시"만 전부 수집
    #    (pic_area 안에 입력일시가 있어도 무시됨)
    # ------------------------------------------------------------
    try:
        photo_els = driver.find_elements(
            By.XPATH,
            "//*[@id='pic_area']"
            "//td[@id='officer_dt' and contains(normalize-space(.),'촬영일시')]"
        )
        for el in photo_els:
            t = _extract_time(el.text.strip())
            if t:
                result["촬영일시목록"].append(t)
    except:
        pass

    # ------------------------------------------------------------
    # 3) GPS 위치확인일시: 기존 방식 유지
    # ------------------------------------------------------------
    try:
        gps = driver.find_element(By.CSS_SELECTOR, "td#gps_dt")
        tx = gps.text.strip()
        result["GPS위치확인일시"] = _extract_time(tx)
    except:
        pass

    return result

def _norm_company_key(s: str) -> str:
    """업소명 매칭용 간단 정규화(공백/법인표기 제거)."""
    if s is None:
        return ""
    t = str(s).strip()
    t = t.replace("(주)", "").replace("㈜", "").replace("주식회사", "")
    t = re.sub(r"\s+", "", t)
    return t

def relax_env_input_time_by_company(sample_rows_map: dict, excel_meta_map: dict):
    """
    동일 날짜 + 동일 업소 방문 케이스에서
    환경기술인 입력일시가 '해당 시료' 채취시간을 벗어나도
    같은 업소의 다른 시료 채취시간 범위 안이면 OK로 완화한다.
    """
    from collections import defaultdict

    windows = defaultdict(list)  # (date, compkey) -> [(start_dt, end_dt), ...]

    for sn, meta in (excel_meta_map or {}).items():
        if not isinstance(meta, dict):
            continue
        date = (meta.get("날짜") or "").strip()
        comp = _norm_company_key(meta.get("업소명", ""))
        st = _pd(meta.get("측정시작DT", ""))
        ed = _pd(meta.get("측정종료DT", ""))
        if date and comp and st and ed:
            windows[(date, comp)].append((st, ed))

    if not windows:
        return

    for sn, rows in (sample_rows_map or {}).items():
        meta = (excel_meta_map or {}).get(sn, {})
        if not isinstance(meta, dict):
            continue

        date = (meta.get("날짜") or "").strip()
        comp = _norm_company_key(meta.get("업소명", ""))
        if not date or not comp:
            continue

        key = (date, comp)
        if key not in windows:
            continue

        for r in rows:
            if not isinstance(r, dict):
                continue
            if r.get("항목") != "환경기술인입력일시":
                continue
            if r.get("비교") != "NG":
                continue

            dt = _pd(r.get("사이트값", ""))
            if not dt:
                continue

            if any(st <= dt <= ed for st, ed in windows[key]):
                r["비교"] = "OK"
                r["사이트만존재"] = ""


def relax_env_input_time_by_env_psic(sample_rows_map: dict, excel_meta_map: dict):
    """
    동일 날짜 + 동일 환경기술인(탭3 field_officer_name) 케이스에서
    환경기술인 입력일시가 '해당 시료' 채취시간을 벗어나도
    같은 환경기술인의 다른 시료 채취시간 범위 안이면 OK로 완화한다.
    """
    from collections import defaultdict

    windows = defaultdict(list)  # (date, psic_name) -> [(start_dt, end_dt), ...]

    for sn, meta in (excel_meta_map or {}).items():
        if not isinstance(meta, dict):
            continue
        date = (meta.get("날짜") or "").strip()
        psic = (meta.get("환경기술인") or "").strip()
        st = _pd(meta.get("측정시작DT", ""))
        ed = _pd(meta.get("측정종료DT", ""))
        if date and psic and st and ed:
            windows[(date, psic)].append((st, ed))

    if not windows:
        return

    for sn, rows in (sample_rows_map or {}).items():
        meta = (excel_meta_map or {}).get(sn, {})
        if not isinstance(meta, dict):
            continue

        date = (meta.get("날짜") or "").strip()
        psic = (meta.get("환경기술인") or "").strip()
        if not date or not psic:
            continue

        key = (date, psic)
        if key not in windows:
            continue

        for r in rows:
            if not isinstance(r, dict):
                continue
            if r.get("항목") != "환경기술인입력일시":
                continue
            if r.get("비교") != "NG":
                continue

            dt = _pd(r.get("사이트값", ""))
            if not dt:
                continue

            if any(st <= dt <= ed for st, ed in windows[key]):
                r["비교"] = "OK"
                r["사이트만존재"] = ""


def _collect_tab1_data(driver, data: dict):
    if not click_tab(driver, "ui-id-1"):
        return
    time.sleep(1.5)
    try:
        els = driver.find_elements(
            By.CSS_SELECTOR,
            "#machineDiv > div > span > span.selection > span > ul > li",
        )
        data["장비"] = [clean_leading_mark(e.text) for e in els if e.text.strip()]
    except Exception:
        data["장비"] = []

    try:
        els = driver.find_elements(
            By.CSS_SELECTOR,
            "#carSection > div > span > span.selection > span > ul > li",
        )
        data["차량"] = [clean_leading_mark(e.text) for e in els if e.text.strip()]
    except Exception:
        data["차량"] = []

    try:
        els = driver.find_elements(
            By.CSS_SELECTOR,
            "#wid-id-4 > div > div.widget-body.no-padding > div > fieldset "
            "> div.row.input-full > section:nth-child(2) "
            "> span > span.selection > span > ul > li",
        )
        data["인력"] = [clean_leading_mark(e.text) for e in els if e.text.strip()]
    except Exception:
        data["인력"] = []

    try:
        els = driver.find_elements(
            By.CSS_SELECTOR,
            "#inairTargetItem > div:nth-child(2) > div > span > span.selection > span > ul > li",
        )
        arr = []
        for e in els:
            t = clean_leading_mark(e.text.strip())
            if t:
                arr.append(t)
        data["측정항목"] = arr
    except Exception:
        data["측정항목"] = []

    try:
        sel_el = driver.find_element(By.ID, "edit_meas_purpose")
        purpose_val = driver.execute_script(
            "var sel = arguments[0]; return sel.options[sel.selectedIndex].text;",
            sel_el,
        )
        data["측정목적"] = purpose_val.strip() if purpose_val else ""
    except Exception:
        data["측정목적"] = ""

    # 측정시설 (Select2 #edit_emis_fac_no) — 표시 텍스트 우선
    try:
        fac_txt = ""
        try:
            fac_txt = driver.find_element(
                By.CSS_SELECTOR, "#select2-edit_emis_fac_no-container"
            ).get_attribute("title") or ""
            fac_txt = (fac_txt or "").strip()
            if not fac_txt:
                fac_txt = driver.find_element(
                    By.CSS_SELECTOR, "#select2-edit_emis_fac_no-container"
                ).text.strip()
        except Exception:
            fac_txt = ""
        if not fac_txt:
            sel_el = driver.find_element(By.CSS_SELECTOR, SEL_EMIS_FAC)
            fac_txt = driver.execute_script(
                """
                var sel = arguments[0];
                if (!sel || sel.selectedIndex < 0) return '';
                return (sel.options[sel.selectedIndex].text || '').trim();
                """,
                sel_el,
            ) or ""
        data["측정시설"] = str(fac_txt).strip()
    except Exception:
        data["측정시설"] = ""

def _collect_tab2_data(driver, data: dict):
    if not click_tab(driver, "ui-id-2"):
        return
    time.sleep(1)
    data["날짜"] = norm_ymd(gv(driver, SEL_DATE))
    data["기상"] = get_weather_text(driver)
    data["기온"] = gv(
        driver,
        "#idWHArea > div > div:nth-child(2) > fieldset > label.col.col-12 "
        "> table > tbody > tr > td:nth-child(2) > input",
    )
    data["습도"] = gv(
        driver,
        "#idWHArea > div > div:nth-child(2) > fieldset > label.col.col-12 "
        "> table > tbody > tr > td:nth-child(3) > input",
    )
    data["기압"] = gv(
        driver,
        "#idWHArea > div > div:nth-child(2) > fieldset > label.col.col-12 "
        "> table > tbody > tr > td:nth-child(4) > input",
    )
    data["풍향"] = get_wind_direction_text(driver)
    data["풍속"] = to_float1(
        gv(
            driver,
            "#idWHArea > div > div:nth-child(2) > fieldset > label.col.col-12 "
            "> table > tbody > tr > td:nth-child(6) > input",
        )
    )
    data["채취시작"] = trim_time_to_hm(gv(driver, SEL_START_TIME))
    data["채취끝"] = trim_time_to_hm(gv(driver, SEL_END_TIME))
    data["표준산소농도"] = to_float1(gv(driver, SEL_O2_STD))
    data["실측산소농도"] = to_float1(gv(driver, SEL_O2_MEAS))
    data["배출가스유량전"] = to_float1(gv(driver, SEL_GAS_VOL_PRE))
    data["배출가스유량후"] = to_float1(gv(driver, SEL_GAS_VOL_POST))
    data["수분량"] = gv(driver, SEL_MOISTURE)
    data["배출가스온도"] = gv(driver, SEL_GAS_TEMP)
    data["배출가스유속"] = to_float2(gv(driver, SEL_GAS_SPEED))


def _collect_tab3_data(driver, data: dict, sample_no: str = "") -> bool:
    if not click_tab(driver, "ui-id-3"):
        return False
    time.sleep(2)
    mob = get_mobile_times(driver)
    data["환경기술인입력일시"] = mob["환경기술인입력일시"]
    data["GPS위치확인일시"] = mob["GPS위치확인일시"]
    data["촬영일시목록"] = mob["촬영일시목록"]
    try:
        data["환경기술인"] = gv(driver, "#field_officer_name")
    except Exception:
        data["환경기술인"] = ""

    sno = (sample_no or data.get("시료번호") or "").strip()
    photos = []
    if sno:
        try:
            photos = download_field_photos(driver, sno, mob.get("촬영일시목록") or [])
        except Exception as e:
            print(f"⚠ 현장사진 저장 실패: {e}")
            photos = []
    data["현장사진"] = photos
    return True


def read_site_data(driver, sample_no):
    """사이트 탭1~3 + PDF 수집. 반환: (data, ok, failures)"""
    data = {}
    failures = []

    _collect_tab1_data(driver, data)
    _collect_tab2_data(driver, data)

    click_tab(driver, "ui-id-2")
    time.sleep(0.5)
    try:
        data["PDF경로"] = download_pdf(driver, sample_no)
    except Exception as e:
        print("⚠ PDF 다운로드 실패:", e)
        data["PDF경로"] = ""
    if not data.get("PDF경로"):
        failures.append("pdf")

    if not _collect_tab3_data(driver, data, sample_no=sample_no):
        failures.append("tab3")

    return data, (len(failures) == 0), failures


# ------------------------------------------------------------
# NAS 검색 / 엑셀 읽기 / 비교 / 저장 (원본 그대로 유지)
# ------------------------------------------------------------
def find_excel_for_sample(sample_no):
    """NAS 에서 sample_no 와 정확 형식 일치 엑셀 파일 검색 (strict)"""
    result = _find_best_file_util(
        sample_no,
        nas_base=NAS_BASE,
        nas_dirs=NAS_DIRS,
        extensions=(".xlsm", ".xlsx", ".xls"),
        strict=True,
    )
    if not result:
        print(" ❌ 엑셀 없음:", sample_no)
    return result


def get_team_no_from_sample(sample_no):
    try:
        return sample_no[7]
    except:
        return ""




def parse_team_input(s: str):
    """팀 입력 문자열을 팀번호 리스트로 파싱.
    허용 예)
      - "" / 공백 : 전체
      - "3"
      - "1,3,5" / "1 3 5"
      - "1-3" (범위)
    반환: ["1","3","5"] 처럼 문자열 리스트(중복 제거, 정렬)
    """
    if s is None:
        return []
    s = str(s).strip()
    if not s:
        return []
    s = s.replace(" ", ",")
    parts = [p.strip() for p in s.split(",") if p.strip()]
    teams = set()
    for p in parts:
        if "-" in p:
            a, b = p.split("-", 1)
            if a.strip().isdigit() and b.strip().isdigit():
                a_i, b_i = int(a), int(b)
                lo, hi = (a_i, b_i) if a_i <= b_i else (b_i, a_i)
                for t in range(lo, hi + 1):
                    if 1 <= t <= 5:
                        teams.add(str(t))
            continue
        if p.isdigit():
            t = int(p)
            if 1 <= t <= 5:
                teams.add(str(t))
    return sorted(teams, key=lambda x: int(x))



# ------------------------------------------------------------
# 비교 관련 로직 (원본 유지)
# ------------------------------------------------------------
SIMPLE_FIELDS = [
    "날짜",
    "기상", "기온", "습도", "기압", "풍향", "풍속",
    "표준산소농도", "실측산소농도",
    "배출가스유량전", "배출가스유량후",
    "수분량", "배출가스온도", "배출가스유속",
    "채취시작", "채취끝",
]




# =====================================================================
# 탭2 RealGrid(측정항목별 테이블) 읽기 + 성적서(엑셀) 기대값 생성 + 비교
#  - eco_check: "수기 입력 오타" 잡는 용도
# =====================================================================

REALGRID_ROOT_CSS = "#measGridAnalySampAnzeDataAirItemList1"

def _rg_norm_date(v):
    if v is None:
        return ""
    s = str(v).strip()
    # '2026-01-02' 형태로 들어오면 그대로, '2026.01.02' 등은 치환
    s = s.replace(".", "-").replace("/", "-")
    # 'YYYY-MM-DD HH:MM' 같이 오면 날짜만
    if " " in s:
        s = s.split(" ")[0]
    return s

def _rg_norm_time(v):
    if v is None:
        return ""
    s = str(v).strip()
    # 'HH:MM:SS' -> 'HH:MM'
    if ":" in s:
        parts = s.split(":")
        if len(parts) >= 2:
            return f"{parts[0]}:{parts[1]}"
    return s

def _rg_norm_num(v):
    if v is None:
        return ""
    s = str(v).strip()
    if s == "":
        return ""
    s = s.replace(",", "")
    try:
        # 1.0 / 1.00 같은 표현을 통일
        f = float(s)
        # 정수면 정수 문자열로
        if abs(f - round(f)) < 1e-9:
            return str(int(round(f)))
        # 소수는 불필요한 0 제거
        out = f"{f:.10f}".rstrip("0").rstrip(".")
        return out
    except:
        return s

def _rg_norm_vol_text(v):
    """시료채취량: 소수점 4자리로 반올림 통일 (사이트·엑셀 비교용)"""
    if v is None:
        return ""
    s = str(v).strip()
    if s == "":
        return ""
    s = s.replace(",", "")
    try:
        return f"{float(s):.4f}"
    except ValueError:
        return s



def _spd_unit_equiv_for_compare(sv: str, ev: str) -> bool:
    """흡인속도 단위 비교 예외: L-MIN 과 L/min 은 동일로 취급(표시는 그대로)."""
    s = "" if sv is None else str(sv).strip()
    e = "" if ev is None else str(ev).strip()
    if not s or not e:
        return False
    # 'L-MIN'만 예외 허용 (대소문자 무시). 다른 변형(L/MIN 등)은 건드리지 않음.
    if s.upper() == "L-MIN" and e.lower() == "l/min":
        return True
    if e.upper() == "L-MIN" and s.lower() == "l/min":
        return True
    return False

def _vol_unit_equiv_for_compare(sv: str, ev: str) -> bool:
    """시료채취량 단위 비교 예외: SM3, Sm3, Sm³ 등을 동일하게 취급."""
    s = ("" if sv is None else str(sv)).strip().lower()
    e = ("" if ev is None else str(ev)).strip().lower()
    if not s or not e:
        return False
    
    # 정규화: ³, ^3 등 모든 변형을 3으로 통일
    s = s.replace("³", "3").replace("^3", "3").replace(" ", "")
    e = e.replace("³", "3").replace("^3", "3").replace(" ", "")
    
    return s == e



def build_excel_realgird_expected(excel_path: str, sample_no: str, is_dust: bool) -> dict:
    """
    성적서 엑셀에서 탭2 RealGrid에 들어가야 하는 값(기대값) 생성.
    - 일반: 입력(분석값) 시트의 헤더(측정시작/측정 종료/시료흡인속도/시료채취량) 기반
    """
    wb = load_workbook(excel_path, data_only=True)

    ws = _find_sheet_by_candidates_openpyxl(wb, ["입력(분석값)", "입력"])
    if ws is None:
        raise RuntimeError("엑셀에서 '입력(분석값)' 시트를 찾지 못함")

    # 1행 헤더 → 열번호 매핑
    header_map = {}
    for col in range(1, ws.max_column + 1):
        v = ws.cell(row=1, column=col).value
        if not v:
            continue
        t = str(v).strip()
        header_map[t] = col

    def col_of(*names):
        for n in names:
            if n in header_map:
                return header_map[n]
        return None

    c_start = col_of("측정시작", "측정 시작")
    c_end   = col_of("측정 종료", "측정종료")
    c_spd   = col_of("시료흡인속도", "시료 흡인속도")
    c_vol   = col_of("시료채취량", "시료 채취량")

    if not (c_start and c_end and c_spd and c_vol):
        raise RuntimeError(f"입력(분석값) 1행 헤더 매칭 실패: start={c_start}, end={c_end}, spd={c_spd}, vol={c_vol} / 헤더={list(header_map.keys())}")

    date_str = sample_to_datestr(sample_no) or ""

    out = {}
    for r in range(2, 65):

        # ✅ A열이 빈칸이면 "측정 안함" → excel_rg에서 제외
        a_flag = ws.cell(row=r, column=1).value  # A열
        if a_flag is None or str(a_flag).strip() == "":
            continue

        item = ws.cell(row=r, column=2).value  # B열(측정항목)
        if not item:
            continue
        item = str(item).strip()
        if not item:
            continue

        st = _rg_norm_time(ws.cell(row=r, column=c_start).value)
        et = _rg_norm_time(ws.cell(row=r, column=c_end).value)
        spd = _rg_norm_num(ws.cell(row=r, column=c_spd).value)
        vol = _rg_norm_vol_text(ws.cell(row=r, column=c_vol).value)

        vol_unit = "Sm³" if item in sm3_items else "L"

        out[item] = {
            "sd": date_str, "st": st,
            "ed": date_str, "et": et,
            "vol": vol, "vol_u": vol_unit,
            "spd": spd, "spd_u": "L/min",
        }

    return out

def build_realgird_compare_rows(sample_no: str, site_rg: dict, excel_rg: dict) -> list:
    """
    RealGrid 비교 결과를 eco_check의 결과 row 포맷으로 반환.
    SKIP_SPEED_ONLY: 흡인속도/단위 비교 PASS
    SKIP_VOL_AND_SPEED: 시료채취량/단위 + 흡인속도/단위 비교 PASS
    """
    rows = []

    site_items = set(site_rg.keys()) if isinstance(site_rg, dict) else set()
    excel_items = set(excel_rg.keys()) if isinstance(excel_rg, dict) else set()
    all_items = sorted(site_items | excel_items)

    fields = [
        ("sd", "측정일(시작)", _rg_norm_date),
        ("st", "시작시간", _rg_norm_time),
        ("ed", "측정일(종료)", _rg_norm_date),
        ("et", "종료시간", _rg_norm_time),
        ("vol", "시료채취량", _rg_norm_vol_text),
        ("vol_u", "채취량단위", lambda x: "" if x is None else str(x).strip()),
        ("spd", "흡인속도", _rg_norm_num),
        ("spd_u", "흡인속도단위", lambda x: "" if x is None else str(x).strip()),
    ]

    for item in all_items:
        # ✅ 항상 먼저 정의 (NameError 방지)
        item_name = ("" if item is None else str(item)).strip()

        s = site_rg.get(item) if item in site_items else None
        e = excel_rg.get(item) if item in excel_items else None

        if s is None:
            rows.append({
                "sample": sample_no,
                "항목": f"[RealGrid] {item_name} (사이트에 없음)",
                "사이트값": "",
                "엑셀값": "존재",
                "비교": "NG",
                "사이트만존재": "",
                "엑셀만존재": "O"
            })
            continue

        if e is None:
            rows.append({
                "sample": sample_no,
                "항목": f"[RealGrid] {item_name} (엑셀에 없음)",
                "사이트값": "존재",
                "엑셀값": "",
                "비교": "NG",
                "사이트만존재": "O",
                "엑셀만존재": ""
            })
            continue

        def is_blank(v):
            return v is None or str(v).strip() == ""

        # ✅ 스킵 규칙
        skip_speed = (item_name in SKIP_SPEED_ONLY)
        skip_vol_and_speed = (item_name in SKIP_VOL_AND_SPEED)

        for key, label, norm_fn in fields:

            # SKIP_SPEED_ONLY → 흡인속도/단위 PASS (단, 사이트값 있으면 NG)
            if skip_speed and key in ("spd", "spd_u"):
                raw = s.get(key) if isinstance(s, dict) else ""
                try:
                    sv = norm_fn(raw)
                except Exception:
                    sv = raw

                if not is_blank(sv):
                    rows.append({
                        "sample": sample_no,
                        "항목": f"[RealGrid] {item_name} / {label} (예외항목: 사이트는 빈칸이어야 함)",
                        "사이트값": sv,
                        "엑셀값": "",
                        "비교": "NG",
                        "사이트만존재": sv,
                        "엑셀만존재": ""
                    })
                continue

            # SKIP_VOL_AND_SPEED → 채취량/단위 + 흡인속도/단위 PASS (단, 사이트값 있으면 NG)
            if skip_vol_and_speed and key in ("vol", "vol_u", "spd", "spd_u"):
                raw = s.get(key) if isinstance(s, dict) else ""
                try:
                    sv = norm_fn(raw)
                except Exception:
                    sv = raw

                if not is_blank(sv):
                    rows.append({
                        "sample": sample_no,
                        "항목": f"[RealGrid] {item_name} / {label} (예외항목: 사이트는 빈칸이어야 함)",
                        "사이트값": sv,
                        "엑셀값": "",
                        "비교": "NG",
                        "사이트만존재": sv,
                        "엑셀만존재": ""
                    })
                continue


            sv = norm_fn(s.get(key)) if isinstance(s, dict) else ""
            ev = norm_fn(e.get(key)) if isinstance(e, dict) else ""
            
            # 단위 예외 처리 (흡인속도단위, 채취량단위)
            ok = (sv == ev)
            if not ok:
                if key == "spd_u" and _spd_unit_equiv_for_compare(sv, ev):
                    ok = True
                elif key == "vol_u" and _vol_unit_equiv_for_compare(sv, ev):
                    ok = True

            rows.append({
                "sample": sample_no,
                "항목": f"[RealGrid] {item_name} / {label}",
                "사이트값": sv,
                "엑셀값": ev,
                "비교": "OK" if ok else "NG",
                "사이트만존재": "",
                "엑셀만존재": ""
            })

    return rows

def compare_scalar(sample, field, site_val, excel_val):
    return {
        "sample": sample,
        "항목": field,
        "사이트값": site_val,
        "엑셀값": excel_val,
        "비교": "OK" if str(site_val) == str(excel_val) else "NG",
        "사이트만존재": "",
        "엑셀만존재": "",
    }


def build_list_status_compare_row(
    sample_no: str,
    result: str,
    status_text: str,
    expected: str = CHECK_LIST_EXPECTED_STATUS,
):
    """
    목록 RealGrid '상태' 열 비교 행.
    - 기대값(엑셀값): '측정분석결과 입력중'
    - 일치 → OK, 다르거나 확인불가 → NG
    """
    site_val = (status_text or "").strip()
    if result == "성공":
        return {
            "sample": sample_no,
            "항목": "목록상태",
            "사이트값": site_val or expected,
            "엑셀값": expected,
            "비교": "OK",
            "사이트만존재": "",
            "엑셀만존재": "",
        }
    if result == "실패":
        return {
            "sample": sample_no,
            "항목": "목록상태",
            "사이트값": site_val or "(상태 불일치)",
            "엑셀값": expected,
            "비교": "NG",
            "사이트만존재": "",
            "엑셀만존재": "",
        }
    return {
        "sample": sample_no,
        "항목": "목록상태",
        "사이트값": site_val or "(확인불가)",
        "엑셀값": expected,
        "비교": "NG",
        "사이트만존재": "",
        "엑셀만존재": "",
    }


def compare_list(sample, field, site_list, excel_list):
    # 탭1 인력·차량·장비: 슬래시 규칙·공백 제거 후 비교 (eco_input 탭1 등록과 동일)
    if field in ("인력", "차량", "장비"):
        site_list = normalize_tab1_select_field(field, site_list)
        excel_list = normalize_tab1_select_field(field, excel_list)
    s = set([x.strip() for x in site_list if x.strip()])
    e = set([x.strip() for x in excel_list if x.strip()])

    only_s = sorted(list(s - e))
    only_e = sorted(list(e - s))

    return {
        "sample": sample,
        "항목": field,
        "사이트값": ", ".join(sorted(s)),
        "엑셀값": ", ".join(sorted(e)),
        "비교": "OK" if not only_s and not only_e else "NG",
        "사이트만존재": ", ".join(only_s),
        "엑셀만존재": ", ".join(only_e),
    }


def _pd(s):
    return parse_datetime_text(s)


def compare_mobile_single(sample, label, t, es, ee):
    if not t:
        return {
            "sample": sample, "항목": label,
            "사이트값": "", "엑셀값": f"{es} ~ {ee}",
            "비교": "NG", "사이트만존재": "", "엑셀만존재": ""
        }

    dt = _pd(t)
    st = _pd(es)
    ed = _pd(ee)
    if not dt or not st or not ed:
        ok = "확인불가"
    else:
        ok = "OK" if st <= dt <= ed else "NG"

    return {
        "sample": sample,
        "항목": label,
        "사이트값": t,
        "엑셀값": f"{es} ~ {ee}",
        "비교": ok,
        "사이트만존재": "" if ok == "OK" else t,
        "엑셀만존재": "",
    }


def compare_mobile_photos(sample, arr, es, ee):
    if not arr:
        return {
            "sample": sample, "항목": "사진촬영일시(전체)",
            "사이트값": "", "엑셀값": f"{es} ~ {ee}",
            "비교": "NG", "사이트만존재": "", "엑셀만존재": ""
        }

    st = _pd(es)
    ed = _pd(ee)
    if not st or not ed:
        return {
            "sample": sample,
            "항목": "사진촬영일시(전체)",
            "사이트값": ", ".join(arr),
            "엑셀값": f"{es} ~ {ee}",
            "비교": "확인불가",
            "사이트만존재": "",
            "엑셀만존재": "",
        }

    out_range = []
    for t in arr:
        dt = _pd(t)
        if not dt or dt < st or dt > ed:
            out_range.append(t)

    return {
        "sample": sample,
        "항목": "사진촬영일시(전체)",
        "사이트값": ", ".join(arr),
        "엑셀값": f"{es} ~ {ee}",
        "비교": "OK" if not out_range else "NG",
        "사이트만존재": ", ".join(out_range),
        "엑셀만존재": "",
    }


def build_comparison_rows(sample_no, site, excel):
    rows = []

    rows.append({
        "sample": sample_no,
        "항목": "엑셀 시료번호 일치 여부",
        "사이트값": sample_no,
        "엑셀값": excel.get("엑셀시료번호", ""),
        "비교": "OK" if sample_no == excel.get("엑셀시료번호", "") else "NG",
        "사이트만존재": "",
        "엑셀만존재": "",
    })

    # --------------------------------------------------
    # 측정목적 체크 (사이트: 자가측정용/참고용 등, 엑셀: 1/2)
    # --------------------------------------------------
    s_purpose = site.get("측정목적", "")
    e_purpose = excel.get("측정목적", "")
    purpose_ok = False
    if s_purpose == "자가측정용" and e_purpose == "1":
        purpose_ok = True
    elif s_purpose == "참고용" and e_purpose == "2":
        purpose_ok = True
    elif not s_purpose and not e_purpose:
        purpose_ok = True
    
    rows.append({
        "sample": sample_no,
        "항목": "측정목적",
        "사이트값": s_purpose,
        "엑셀값": e_purpose,
        "비교": "OK" if purpose_ok else "NG",
        "사이트만존재": "",
        "엑셀만존재": "",
    })

    # 측정시설 ↔ 입력!E4 측정인 시설명
    s_fac = site.get("측정시설", "")
    e_fac = excel.get("측정인시설명", "")
    fac_ok = facility_labels_match(s_fac, e_fac)
    rows.append({
        "sample": sample_no,
        "항목": "측정시설",
        "사이트값": s_fac,
        "엑셀값": e_fac,
        "비교": "OK" if fac_ok else "NG",
        "사이트만존재": "",
        "엑셀만존재": "",
    })

    # ★ 비산먼지면 특정 필드는 비교 PASS
    is_dust = bool(excel.get("is_dust") or site.get("is_dust"))

    for f in SIMPLE_FIELDS:
        if is_dust and f in DUST_SKIP_FIELDS:
            continue
        rows.append(compare_scalar(sample_no, f, site.get(f, ""), excel.get(f, "")))

    rows.append(compare_list(sample_no, "측정항목", site.get("측정항목", []), excel.get("측정항목", [])))
    rows.append(compare_list(sample_no, "장비", site.get("장비", []), excel.get("장비", [])))
    rows.append(compare_list(sample_no, "차량", site.get("차량", []), excel.get("차량", [])))
    rows.append(compare_list(sample_no, "인력", site.get("인력", []), excel.get("인력", [])))

    es = excel.get("측정시작DT", "")
    ee = excel.get("측정종료DT", "")

    rows.append({
        "sample": sample_no,
        "항목": "환경기술인",
        "사이트값": site.get("환경기술인", ""),
        "엑셀값": "",
        "비교": "",
        "사이트만존재": "",
        "엑셀만존재": "",
    })
    rows.append(compare_mobile_single(sample_no, "환경기술인입력일시",
                                      site.get("환경기술인입력일시", ""), es, ee))
    rows.append(compare_mobile_single(sample_no, "GPS위치확인일시",
                                      site.get("GPS위치확인일시", ""), es, ee))
    rows.append(compare_mobile_photos(sample_no, site.get("촬영일시목록", []), es, ee))

    # 현장사진 1~3 — 촬영일시만 (경로는 미리보기로 확인)
    photos = site.get("현장사진") or []
    if not isinstance(photos, list):
        photos = []
    for i in range(1, 4):
        p = next((x for x in photos if isinstance(x, dict) and int(x.get("idx") or 0) == i), None)
        if p is None and i - 1 < len(photos) and isinstance(photos[i - 1], dict):
            p = photos[i - 1]
        path = (p or {}).get("path", "") if p else ""
        shot = (p or {}).get("shot_at", "") if p else ""
        if not shot:
            arr = site.get("촬영일시목록") or []
            if i - 1 < len(arr):
                shot = arr[i - 1]
        has_file = bool(path and os.path.isfile(path))
        rows.append({
            "sample": sample_no,
            "항목": f"현장사진{i}",
            "사이트값": shot,
            "엑셀값": "",
            "비교": "OK" if has_file else "NG",
            "사이트만존재": "" if has_file else (shot or "사진파일없음"),
            "엑셀만존재": "",
        })

    # RealGrid 비교는 그대로 (아래 2)에서 로직 추가)
    try:
        site_rg = site.get("realgrid", {}) if isinstance(site, dict) else {}
        excel_rg = excel.get("realgrid", {}) if isinstance(excel, dict) else {}
        if site_rg or excel_rg:
            rows.extend(build_realgird_compare_rows(sample_no, site_rg, excel_rg))
    except Exception as e:
        rows.append({
            "시료번호": sample_no,
            "항목": "[RealGrid] 비교 중 예외",
            "사이트값": str(e),
            "엑셀값": "",
            "비교": "NG",
            "사이트만존재": "",
            "엑셀만존재": ""
        })

    return rows
          

# ------------------------------------------------------------
# 결과 저장
# ------------------------------------------------------------
def _next_available_path(path: str) -> str:
    base, ext = os.path.splitext(path)
    for i in range(1, 1000):
        cand = f"{base}_{i}{ext}"
        if not os.path.exists(cand):
            return cand
    return f"{base}_{int(time.time())}{ext}"


def _photo_paths_from_map(sample_no: str) -> list[str]:
    photos = PHOTO_MAP.get(sample_no) or []
    paths = []
    for p in photos:
        if isinstance(p, dict):
            path = (p.get("path") or "").strip()
        else:
            path = str(p or "").strip()
        if path and os.path.isfile(path):
            paths.append(path)
    return paths


def _photo_shots_from_map(sample_no: str) -> list[str]:
    photos = PHOTO_MAP.get(sample_no) or []
    shots = []
    for p in photos:
        if isinstance(p, dict):
            shots.append((p.get("shot_at") or "").strip())
        else:
            shots.append("")
    return shots


def _embed_field_photos(
    ws,
    photo_paths: list[str],
    *,
    start_col: int = 7,
    anchor_row: int | None = None,
    title: str = "현장사진 미리보기",
    shot_times: list[str] | None = None,
    max_w: int = 280,
    max_h: int = 210,
    row_height: float = 160,
    labels_on_anchor_row: bool = False,
) -> int:
    """시트에 현장사진 미리보기 삽입. 다음 사용 가능한 행 번호 반환.

    labels_on_anchor_row=True 이면
      anchor_row: (시료번호 등과 같은 줄) PIC 라벨
      anchor_row+1: 이미지
    기본(False)은 title 행 → 라벨 → 이미지.
    """
    try:
        from openpyxl.drawing.image import Image as XLImage
        from openpyxl.utils import get_column_letter
    except Exception as e:
        print(f"⚠ 이미지 삽입 모듈 실패: {e}")
        return anchor_row or (ws.max_row + 2)

    if not photo_paths:
        return anchor_row or (ws.max_row + 2)

    row0 = anchor_row if anchor_row is not None else max(ws.max_row + 2, 2)
    if labels_on_anchor_row:
        label_row = row0
        img_row = row0 + 1
    else:
        if title:
            ws.cell(row=row0, column=start_col, value=title)
            label_row = row0 + 1
            img_row = row0 + 2
        else:
            label_row = row0
            img_row = row0 + 1

    col = start_col
    shots = list(shot_times or [])
    for i, path in enumerate(photo_paths, start=1):
        if not path or not os.path.isfile(path):
            continue
        try:
            label = f"PIC{i}"
            if i - 1 < len(shots) and shots[i - 1]:
                label = f"PIC{i} ({shots[i - 1]})"
            ws.cell(row=label_row, column=col, value=label)
            img = XLImage(path)
            try:
                ow, oh = float(img.width or max_w), float(img.height or max_h)
                scale = min(max_w / ow, max_h / oh, 1.0)
                img.width = int(ow * scale)
                img.height = int(oh * scale)
            except Exception:
                img.width = max_w
                img.height = max_h
            img.anchor = f"{get_column_letter(col)}{img_row}"
            ws.add_image(img)
            ws.column_dimensions[get_column_letter(col)].width = max(38, int(max_w / 7))
            ws.row_dimensions[img_row].height = row_height
            col += 2
        except Exception as e:
            print(f"⚠ 현장사진 삽입 실패({path}): {e}")
    return img_row + 2


def save_results(sample_rows_map, out_path):
    wb = Workbook()
    ws_sum = wb.active
    ws_sum.title = "요약"
    ws_sum.append(["시료번호", "업소명", "항목", "비교", "사이트값", "엑셀값", "사이트만존재", "엑셀만존재"])
    # 요약 시트 1행 1열 틀고정 (B2 셀 기준)
    ws_sum.freeze_panes = "B2"

    for sample_no, rows in sample_rows_map.items():
        ws = wb.create_sheet(sample_no)
        ws.append(["항목", "사이트값", "엑셀값", "비교", "사이트만존재", "엑셀만존재"])
        # ✅ 개별 시료 시트 1행 1열 틀고정 (B2 셀 기준)
        ws.freeze_panes = "B2"
        for r in rows:
            ws.append([r["항목"], r["사이트값"], r["엑셀값"], r["비교"],
                       r["사이트만존재"], r["엑셀만존재"]])

            company = COMPANY_MAP.get(sample_no, "")
            ws_sum.append([
                sample_no, company,
                r["항목"], r["비교"],
                r["사이트값"], r["엑셀값"],
                r["사이트만존재"], r["엑셀만존재"]
            ])

        # PDF 하이퍼링크 행 추가
        pdf_path = PDF_MAP.get(sample_no, "")
        sum_pdf_row = None
        if pdf_path:
            row_idx = ws.max_row + 1
            ws.cell(row=row_idx, column=1, value="PDF 열기")
            link_cell = ws.cell(row=row_idx, column=2, value=pdf_path)
            link_cell.hyperlink = pdf_path
            link_cell.style = "Hyperlink"

            sum_pdf_row = ws_sum.max_row + 1
            company = COMPANY_MAP.get(sample_no, "")
            ws_sum.cell(row=sum_pdf_row, column=1, value=sample_no)
            ws_sum.cell(row=sum_pdf_row, column=2, value=company)
            ws_sum.cell(row=sum_pdf_row, column=3, value="PDF 열기")
            ws_sum.cell(row=sum_pdf_row, column=4, value="OK")
            sum_link_cell = ws_sum.cell(row=sum_pdf_row, column=5, value=pdf_path)
            sum_link_cell.hyperlink = pdf_path
            sum_link_cell.style = "Hyperlink"

        # 시료 시트 + 요약(PDF 바로 아래) 미리보기 — sample_rows_map(=해당 팀)만
        photo_paths = _photo_paths_from_map(sample_no)
        if photo_paths:
            _embed_field_photos(
                ws,
                photo_paths,
                shot_times=_photo_shots_from_map(sample_no),
            )

            company = COMPANY_MAP.get(sample_no, "")
            photo_row = (sum_pdf_row + 1) if sum_pdf_row else (ws_sum.max_row + 1)
            ws_sum.cell(row=photo_row, column=1, value=sample_no)
            ws_sum.cell(row=photo_row, column=2, value=company)
            ws_sum.cell(row=photo_row, column=3, value="현장사진")
            ws_sum.cell(row=photo_row, column=4, value="OK")
            next_row = _embed_field_photos(
                ws_sum,
                photo_paths,
                start_col=5,
                anchor_row=photo_row,
                title="",
                shot_times=_photo_shots_from_map(sample_no),
                max_w=220,
                max_h=165,
                row_height=130,
                labels_on_anchor_row=True,
            )
            # 이미지 행 아래로 다음 시료 데이터 밀기
            if next_row and ws_sum.max_row < next_row:
                ws_sum.cell(row=next_row, column=1, value="")

    # NG 빨간색 조건부서식
    red_fill = PatternFill(start_color="FFFF0000", end_color="FFFF0000", fill_type="solid")

    if ws_sum.max_row > 1:
        rule_sum = FormulaRule(formula=['$D2="NG"'], fill=red_fill)
        ws_sum.conditional_formatting.add(f"D2:D{ws_sum.max_row}", rule_sum)
    # 각 시료 시트(D열) - 요약/중복검사 시트들은 제외
    for name in wb.sheetnames:
        if name == "요약":
            continue

        ws_sample = wb[name]
        if ws_sample.max_row > 1:
            rule_sample = FormulaRule(formula=['$D2="NG"'], fill=red_fill)
            ws_sample.conditional_formatting.add(f"D2:D{ws_sample.max_row}", rule_sample)

    # 요약 시트 A~C 열 너비 자동 조정
    autofit_columns(ws_sum, "A:C")

    try:
        wb.save(out_path)
        print(f"\n[완료] 결과 저장 → {out_path}")
    except PermissionError as e:
        alt = _next_available_path(out_path)
        wb.save(alt)
        print(f"\n⚠️ 저장 실패(잠금/권한): {out_path}")
        print(f"[대체 저장] → {alt}")



# ------------------------------------------------------------
# 메인 실행
# ------------------------------------------------------------
def main(progress_callback=None, cancel_event=None):
    print("=== 측정인.kr 자동 비교 시스템 시작 ===")

    login_id = input("측정인아이디: ").strip()
    login_pw = input("측정인비밀번호: ").strip()
    start_date = input("시작일 (YYYY-MM-DD): ").strip()
    end_date = input("종료일 (YYYY-MM-DD): ").strip()
    team_input = input("팀번호(1-5, 예: 3 / 1,3,5 / 1-3, 엔터=전체): ").strip()

    yyyymmdd = start_date.replace("-", "")

    # ------------------------------------------------------------
    # 0) 파일 기반 시료번호 목록 먼저 생성 (팀 필터 포함)
    #    - 선택한 팀이 파일 목록에 없으면: 사이트 로그인/브라우저 실행 없이 종료
    # ------------------------------------------------------------
    samples = collect_samples_from_files(start_date, nas_base=NAS_BASE, nas_dirs=NAS_DIRS)
    print(f"[파일 기반 시료번호] {len(samples)}개 → {samples}")

    teams = parse_team_input(team_input)

    if teams:
        team_set = set(teams)
        samples = [sn for sn in samples if get_team_no_from_sample(sn) in team_set]

        # 필터 후 실제 남은 팀만 추출
        run_teams = sorted(
            {get_team_no_from_sample(sn) for sn in samples if str(get_team_no_from_sample(sn)).isdigit()},
            key=lambda x: int(x)
        )

        print(f"[팀 선택] 실행 팀 → {','.join(run_teams) if run_teams else '-'}팀  |  {len(samples)}개")

    else:
        teams = sorted(
            {get_team_no_from_sample(sn) for sn in samples if str(get_team_no_from_sample(sn)).isdigit()},
            key=lambda x: int(x)
        )
        print(f"[전체 선택] 발견된 팀 자동 분리 대상 → {','.join(teams) if teams else '-'}팀  |  {len(samples)}개")

    if not samples:
        if team_input.strip():
            print(f"❌ 선택한 팀({team_input})에 해당하는 시료번호가 파일에서 발견되지 않음 → 사이트 로그인 없이 종료")
        else:
            print("❌ 해당 날짜 파일에서 시료번호 없음 → 사이트 로그인 없이 종료")
        return

    total_samples = len(samples)
    current_count = 0

    # 단일 팀이면 기존과 동일하게 team_tag 사용(파일명/로그용)
    team_tag = teams[0] if len(teams) == 1 else ""

    out_dir = MEASIN_REVIEW
    out_name = f"{yyyymmdd} 팀{team_tag} 검토 파일.xlsx" if team_tag else f"{yyyymmdd} 검토 파일.xlsx"
    RESULT_XLSX = os.path.join(out_dir, out_name)

    driver = None
    sample_rows = {}
    excel_meta_map = {}  # 시료별 (날짜/업소/채취시간) 메타

    try:
        if is_cancelled(cancel_event):
            print("\n[취소됨] 작업을 중단합니다.")
            return

        driver = init_driver()
        if is_cancelled(cancel_event):
            print("\n[취소됨] 작업을 중단합니다.")
            try:
                driver.quit()
            except Exception:
                pass
            return

        login(driver, login_id, login_pw)
        search_date(driver, start_date, end_date)

        for sample_no in samples:
            if is_cancelled(cancel_event):
                print("\n[취소됨] 작업을 중단합니다.")
                break
                
            current_count += 1
            if progress_callback:
                progress_callback(current_count, total_samples)
                
            print("\n-------------------------------------------")
            print(f" 시료({current_count}/{total_samples}):", sample_no)
            print("-------------------------------------------")

            if is_cancelled(cancel_event):
                print("\n[취소됨] 작업을 중단합니다.")
                break

            sample_done = False
            session_recover_count = 0
            sample_attempt = 0
            while not sample_done:
                if is_cancelled(cancel_event):
                    break

                sample_attempt += 1
                if sample_attempt > MAX_SAMPLE_DETAIL_RETRY:
                    print(f"❌ {sample_no} 재시도 한도 초과 → 다음 시료로")
                    break

                if not ensure_logged_in_or_recover(
                    driver,
                    login_id,
                    login_pw,
                    start_date=start_date,
                    end_date=end_date,
                ):
                    print(f"❌ 로그인 복구 실패 → 다음 시료로 ({sample_no})")
                    break

                if sample_attempt == 1:
                    opened = open_detail_with_session_recovery(
                        driver,
                        sample_no,
                        login_id,
                        login_pw,
                        start_date=start_date,
                        end_date=end_date,
                    )
                else:
                    opened = reopen_sample_from_search(
                        driver,
                        sample_no,
                        login_id=login_id,
                        login_pw=login_pw,
                        start_date=start_date,
                        end_date=end_date,
                    )

                if not opened:
                    if sample_attempt < MAX_SAMPLE_DETAIL_RETRY:
                        continue
                    print(f"❌ 상세페이지 실패 → 다음 시료로 ({sample_no})")
                    break

                if not ensure_detail_page_for_tab1(driver):
                    if sample_attempt < MAX_SAMPLE_DETAIL_RETRY:
                        go_back_to_list(driver)
                        continue
                    print(f"❌ 상세 페이지 확인 실패 → 다음 시료로 ({sample_no})")
                    break

                if is_cancelled(cancel_event):
                    break

                try:
                    time.sleep(1)
                    site, read_ok, read_failures = read_site_data(driver, sample_no)
                    if not read_ok:
                        if sample_attempt < MAX_SAMPLE_DETAIL_RETRY:
                            go_back_to_list(driver)
                            continue

                    PDF_MAP[sample_no] = site.get("PDF경로", "")
                    PHOTO_MAP[sample_no] = site.get("현장사진") or []
                    xlsx = find_excel_for_sample(sample_no)
                    if not xlsx:
                        go_back_to_list(driver)
                        sample_done = True
                        continue
                    is_dust = is_fugitive_dust_file(str(xlsx))
                    excel = parse_measuring_record(str(xlsx), sample_no)
                    COMPANY_MAP[sample_no] = excel.get("업소명", "")
                    # --------------------------------------------------
                    # 시료별 메타 저장(날짜/업소/채취시간) - 환경기술인 입력일시 완화용
                    # --------------------------------------------------
                    excel_meta_map[sample_no] = {
                        "날짜": excel.get("날짜", ""),
                        "업소명": excel.get("업소명", ""),
                        "측정시작DT": excel.get("측정시작DT", ""),
                        "측정종료DT": excel.get("측정종료DT", ""),
                        "환경기술인": site.get("환경기술인", ""),
                    }

                    excel["is_dust"] = is_dust
                    site["is_dust"] = is_dust
                    # --------------------------------------------------
                    # 탭2 RealGrid(항목별 표)도 같이 비교: 수기 입력 오타 탐지용
                    # --------------------------------------------------
                    time.sleep(1)

                    try:
                        site_rg = rg_api_read_data(driver, REALGRID_ROOT_CSS)
                    except Exception as e:
                        site_rg = {}
                        # RealGrid를 못 읽어도 전체 비교는 진행
                        print("⚠ 탭2 RealGrid 읽기 실패:", e)

                    try:
                        excel_rg = build_excel_realgird_expected(str(xlsx), sample_no, is_dust=is_dust)
                    except Exception as e:
                        excel_rg = {}
                        print("⚠ 엑셀 RealGrid 기대값 생성 실패:", e)

                    site["realgrid"] = site_rg
                    excel["realgrid"] = excel_rg

                    rows = build_comparison_rows(sample_no, site, excel)
                    go_back_to_list(driver)

                    # 목록 RealGrid '상태' 열: 측정분석결과 입력중 인지 확인 → 결과엑셀
                    try:
                        st_result, st_text = verify_tab4_list_status(
                            driver,
                            sample_no,
                            success_text=CHECK_LIST_EXPECTED_STATUS,
                        )
                    except Exception as e:
                        st_result, st_text = "확인불가", f"(예외: {e})"
                    rows.append(
                        build_list_status_compare_row(sample_no, st_result, st_text)
                    )
                    if st_result == "성공":
                        print(f"  ✅ 목록상태 OK: {sample_no} ({st_text or CHECK_LIST_EXPECTED_STATUS})")
                    else:
                        print(
                            f"  ❌ 목록상태 NG: {sample_no} "
                            f"→ 사이트='{st_text}' / 기대='{CHECK_LIST_EXPECTED_STATUS}'"
                        )

                    sample_rows[sample_no] = rows
                    relax_env_input_time_by_company(sample_rows, excel_meta_map)
                    relax_env_input_time_by_env_psic(sample_rows, excel_meta_map)
                    if read_ok:
                        print(" 완료 : ", sample_no)
                    else:
                        print(
                            f" ⚠ 부분 완료 : {sample_no} "
                            f"(미수집: {', '.join(read_failures)})"
                        )
                    sample_done = True
                except Exception as e:
                    if is_logged_out(driver) and session_recover_count < 5:
                        session_recover_count += 1
                        print(
                            f"▶ 로그아웃 감지 → 재로그인 후 {sample_no} 재시도 "
                            f"({session_recover_count}/5)"
                        )
                        recover_site_session(
                            driver,
                            login_id,
                            login_pw,
                            start_date=start_date,
                            end_date=end_date,
                        )
                        sample_attempt -= 1
                        continue

                    print(f"❌ 시료 처리 실패 ({sample_no}): {e}")
                    if sample_attempt < MAX_SAMPLE_DETAIL_RETRY:
                        try:
                            go_back_to_list(driver)
                        except Exception:
                            pass
                        continue

                    try:
                        go_back_to_list(driver)
                    except Exception:
                        pass
                    break

        if sample_rows and not is_cancelled(cancel_event):
            # teams가 여러 개인 경우 → 팀별로 파일을 따로 저장 (dash에서 여러 파일 선택해서 종합검토)
            if teams and len(teams) > 1:
                base_dir = os.path.dirname(RESULT_XLSX)
                for t in teams:
                    team_map = {sn: rows for sn, rows in sample_rows.items() if get_team_no_from_sample(sn) == t}
                    if not team_map:
                        continue
                    out_name = f"{yyyymmdd} 팀{t} 검토 파일.xlsx"
                    out_path = os.path.join(base_dir, out_name)
                    save_results(team_map, out_path)
                    print(f"✅ 저장 완료: {out_path}")
            else:
                save_results(sample_rows, RESULT_XLSX)
                print(f"✅ 저장 완료: {RESULT_XLSX}")
        else:
            if is_cancelled(cancel_event):
                print("⚠ 취소되어 결과가 저장되지 않았습니다.")
            else:
                print("⚠ 저장할 결과 없음")

    finally:
        if is_cancelled(cancel_event) and driver:
            try:
                driver.quit()
                print("취소됨: 브라우저를 닫았습니다.")
            except:
                pass
        else:
            print("작업 종료. 브라우저는 직접 닫아도 됨.")


if __name__ == "__main__":
    try:
        from log_utils import run_log
        with run_log("eco_check"):
            main()
    except Exception as e:
        log_error("eco_check.main", e)
        raise
