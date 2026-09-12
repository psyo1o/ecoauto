# -*- coding: utf-8 -*-
"""
그룹웨어(FastAPI) 연동 — 탭4 입력완료 후 측정 데이터·PDF 전송, 로컬 정리 엑셀.

API (그룹웨어 측 구현 완료 가정):
  POST /api/external/report_data   — JSON
  POST /api/reports/sync           — multipart: data(JSON) + file(PDF)  [권장]
  POST /api/external/report_pdf    — PDF만 (report_data 선행 시)
"""
from __future__ import annotations

import json
import os
import re
import time
import warnings
from datetime import datetime

warnings.filterwarnings("ignore", category=UserWarning, module=r"openpyxl")

import requests
from openpyxl import Workbook, load_workbook

from config import (
    GROUPWARE_API_TOKEN,
    GROUPWARE_BASE_URL,
    GROUPWARE_ENABLED,
    GROUPWARE_EXCEL_OUTPUT_DIR,
    GROUPWARE_MAX_RETRIES,
    GROUPWARE_REPORT_DATA_PATH,
    GROUPWARE_REPORT_PDF_PATH,
    GROUPWARE_REPORT_SYNC_PATH,
    GROUPWARE_TIMEOUT_SEC,
)
from data_utils import normalize_company, sample_to_datestr
from excel_utils import autofit_columns, parse_measuring_record
from format_utils import format_time
from log_utils import log_error, log_message
from tab4_utils import _norm_rg

# 전송로그 재전송 상태 (기간 재전송 GUI용)
RESEND_STATUS_PENDING = "대기"
RESEND_STATUS_DONE = "완료"
RESEND_STATUS_SKIP = "스킵"
_RESEND_DONE_STATUSES = {RESEND_STATUS_DONE, RESEND_STATUS_SKIP}


def _clean_cell(v) -> str:
    if v is None:
        return ""
    return str(v).strip()


def _normalize_match_key(value) -> str:
    """
    시설 매칭키 정규화 (E2/E4, 그룹웨어 DB와 동일 규칙으로 비교).
    - 앞뒤·내부 공백 제거, nbsp 제거
    - 시설 보조키(E4) 등 '문의 휴 게소' → '문의휴게소' 형태로 맞춤
    """
    if value is None:
        return ""
    s = str(value).replace("\xa0", " ").strip()
    if not s:
        return ""
    return re.sub(r"\s+", "", s)


def read_input_sheet_keys(excel_path: str) -> dict:
    """
    **대기 전용** — 성적서 「입력」 시트 거래처·시설 매칭키.

    수신 기준: \\\\192.168.10.163\\docker\\approval_mvp\\docs\\ECO_INPUT_MATCHING.md
    그룹웨어 거래처 = 사업장(G열). company_name 은 공장 구분된 사업장명(G 전체).
    의뢰기관(C열)만 보내면 안 됨.

    | 성적서 | 필드 | 그룹웨어 대응 |
    |--------|------|---------------|
    | H7 (라벨 G7 업체명) | company_name | 사업장(G열) |
    | H6 (구 K7/J7 근처) | site_no | 사업장관리번호(H열) — 매칭 최우선 |
    | E2 | sems_stack / sems_serial | 시설 폴백(sems_stack) |
    | E4 (라벨 D4 측정인 시설명) | facility_name | 측정시설명(K열) — 시설 1차 |
    """
    out = {
        "company_name": "",
        "sems_serial": "",
        "facility_alt": "",
        "facility_name": "",
        "biz_no": "",
        "site_no": "",
    }
    try:
        wb = load_workbook(excel_path, data_only=True, read_only=True)
    except Exception:
        return out

    try:
        ws = None
        if "입력" in wb.sheetnames:
            ws = wb["입력"]
        else:
            for name in wb.sheetnames:
                if name.replace(" ", "") == "입력":
                    ws = wb[name]
                    break
        if ws is None:
            return out

        # 사업장명: H7 값 (G7은 보통 '업체명' 라벨)
        company = _clean_cell(ws["H7"].value)
        g7 = _clean_cell(ws["G7"].value)
        if not company and g7 and g7 not in ("업체명", "업소명", "사업장명"):
            company = g7
        out["company_name"] = company

        out["sems_serial"] = _clean_cell(ws["E2"].value)
        fac = _clean_cell(ws["E4"].value)
        out["facility_alt"] = fac
        out["facility_name"] = fac
        out["biz_no"] = ""

        # 사업장관리번호: 주 셀은 입력!H6. 구버전 폴백 K7, L7, 라벨 인접
        site_no = _clean_cell(ws["H6"].value)
        if not site_no:
            site_no = _clean_cell(ws["K7"].value)
        if not site_no:
            site_no = _clean_cell(ws["L7"].value)
        if not site_no:
            for label_cell, value_cell in (
                ("G6", "H6"),
                ("J7", "K7"),
                ("J7", "L7"),
                ("I7", "J7"),
                ("L7", "M7"),
            ):
                if _clean_cell(ws[label_cell].value) == "사업장관리번호":
                    site_no = _clean_cell(ws[value_cell].value)
                    if site_no:
                        break
            if not site_no:
                # row 6 then row 7 near label
                for row in (6, 7):
                    for col in range(7, 15):  # G..N
                        cell = ws.cell(row=row, column=col)
                        if _clean_cell(cell.value) == "사업장관리번호":
                            site_no = _clean_cell(ws.cell(row=row, column=col + 1).value)
                            if site_no:
                                break
                    if site_no:
                        break
        out["site_no"] = site_no

    except Exception as e:
        log_error("groupware.read_input_sheet_keys", e)
    finally:
        try:
            wb.close()
        except Exception:
            pass

    if not out["company_name"]:
        try:
            wb2 = load_workbook(excel_path, data_only=True, read_only=True)
            try:
                if "대기측정기록부" in wb2.sheetnames:
                    v = _clean_cell(wb2["대기측정기록부"]["D3"].value)
                    if v:
                        out["company_name"] = v
            finally:
                wb2.close()
        except Exception:
            pass
    return out


def facility_match_keys(keys: dict) -> dict:
    """API payload용 시설·거래처 매칭 필드.

    수신측(ECO_INPUT_MATCHING.md): facility_name(마스터 K) 우선,
    실패·공란 시 sems_stack / measurements[].stack / 파일명 NO.n 폴백.
    E1(굴뚝번호) 필드로 보내지 않음 — E2는 sems_stack·sems_serial 로 전달.
    """
    sems_serial = _normalize_match_key(keys.get("sems_serial"))
    facility_alt = _normalize_match_key(keys.get("facility_alt"))
    company_raw = _clean_cell(keys.get("company_name"))
    # 측정인 시설명(E4)=그룹웨어 K열. E2 숫자코드로 facility_name 대체하지 않음.
    facility_name = _clean_cell(keys.get("facility_name")) or facility_alt
    biz_no = _clean_cell(keys.get("biz_no"))
    site_no = _clean_cell(keys.get("site_no"))
    return {
        "company_name": company_raw,
        "company_name_norm": normalize_company(company_raw),
        "biz_no": biz_no,
        "site_no": site_no,
        # 수신 extract_outlet_hint 가 읽는 키 (sems_serial 단독은 폴백 미사용)
        "sems_stack": sems_serial,
        "sems_serial": sems_serial,
        "SEMS일련번호": sems_serial,
        "facility_alt": facility_alt,
        "facility_match_order": ["facility_name", "sems_stack"],
        "facility_name": facility_name,
    }


# ---------------------------------------------------------------------------
# 세션 누적 (로컬 정리 엑셀용)
# ---------------------------------------------------------------------------
class GroupwareRunLog:
    """한 번의 자동입력 실행 동안 그룹웨어 전송 기록 + 결과 추적."""

    def __init__(self):
        self.rows: list[dict] = []
        # 시료별 전송 결과: {sample_no: {"data_ok", "pdf_ok", "pdf_attempted", "error", "payload"}}
        self._results: dict[str, dict] = {}

    def add_rows(self, rows: list[dict]):
        self.rows.extend(rows)

    def record_result(self, sample_no: str, result: dict, payload: dict | None = None):
        """전송 결과 기록 (재시도 시 덮어씀)."""
        self._results[sample_no] = {
            "data_ok": result.get("data_ok", False),
            "pdf_ok": result.get("pdf_ok", False),
            "pdf_attempted": result.get("pdf_attempted", False),
            "error": result.get("error", ""),
            "warnings": list(result.get("warnings") or []),
            "verify_key": result.get("verify_key", ""),
            "status": result.get("status", ""),
            "payload": payload or result.get("_payload") or {},
        }

    @property
    def count(self) -> int:
        return len(self.rows)

    def failed_samples(self) -> list[dict]:
        """데이터/PDF 실패 또는 collected_at·measure_date 등 hard 필드 경고가 있는 시료.
        시설명 매칭 soft 경고만 있는 경우는 제외 (데이터는 저장됨).
        """
        out = []
        for sno, r in self._results.items():
            data_fail = not r["data_ok"]
            pdf_fail = r["pdf_attempted"] and not r["pdf_ok"]
            warns = r.get("warnings") or []
            field_warn = _has_critical_field_warnings(warns)
            if data_fail or pdf_fail or field_warn:
                out.append({
                    "sample_no": sno,
                    "data_ok": r["data_ok"],
                    "pdf_ok": r["pdf_ok"],
                    "pdf_attempted": r["pdf_attempted"],
                    "error": r["error"],
                    "warnings": warns,
                    "verify_key": r.get("verify_key", ""),
                    "payload": r["payload"],
                })
        return out

    def facility_soft_samples(self) -> list[dict]:
        """데이터·PDF는 OK, 시설명 매칭 soft 경고만 있는 시료."""
        failed_nos = {f["sample_no"] for f in self.failed_samples()}
        out = []
        for sno, r in self._results.items():
            if sno in failed_nos:
                continue
            soft = _soft_facility_warnings(r.get("warnings") or [])
            if soft and r["data_ok"] and (r["pdf_ok"] or not r["pdf_attempted"]):
                out.append({
                    "sample_no": sno,
                    "warnings": soft,
                    "verify_key": r.get("verify_key", ""),
                    "payload": r["payload"],
                })
        return out

    def summary_text(self) -> str:
        """전송 결과 요약 문자열."""
        total = len(self._results)
        if total == 0:
            return "(그룹웨어 전송 없음)"
        ok = sum(
            1 for r in self._results.values()
            if r["data_ok"]
            and (r["pdf_ok"] or not r["pdf_attempted"])
            and not _has_critical_field_warnings(r.get("warnings") or [])
        )
        fail = total - ok
        lines = [f"그룹웨어 전송 결과: 성공 {ok} / 실패·경고 {fail} / 전체 {total}"]
        for sno, r in self._results.items():
            d = "✅" if r["data_ok"] else "❌"
            p = ("✅" if r["pdf_ok"] else ("－" if not r["pdf_attempted"] else "❌"))
            err = f"  ({r['error'][:60]})" if r["error"] else ""
            warns = r.get("warnings") or []
            wtxt = f"  ⚠{warns}" if warns else ""
            vk = f"  verify_key={r.get('verify_key')}" if r.get("verify_key") else ""
            lines.append(f"  {d}데이터 {p}PDF  {sno}{err}{wtxt}{vk}")
        return "\n".join(lines)


# ---------------------------------------------------------------------------
# 설정
# ---------------------------------------------------------------------------
def is_groupware_enabled(override: bool | None = None) -> bool:
    """override가 있으면 그 값, 없으면 config.ini GROUPWARE.ENABLED (기본 ON)."""
    if override is not None:
        return bool(override)
    return GROUPWARE_ENABLED


def _auth_headers() -> dict:
    h = {"Accept": "application/json"}
    token = (GROUPWARE_API_TOKEN or "").strip()
    if token:
        h["Authorization"] = f"Bearer {token}"
    return h


def _url(path: str) -> str:
    base = (GROUPWARE_BASE_URL or "").rstrip("/")
    p = path if path.startswith("/") else f"/{path}"
    return f"{base}{p}"


# ---------------------------------------------------------------------------
# 추출 유틸
# ---------------------------------------------------------------------------
def format_collected_at(start_hm: str, end_hm: str) -> str:
    """채취시간 `13:10~15:30` 형식 (그룹웨어 DB·진위확인 권장)."""
    s = format_time(start_hm) if start_hm else ""
    e = format_time(end_hm) if end_hm else ""
    if s and e:
        return f"{s}~{e}"
    return s or e or ""


def _normalize_measure_date(raw, sample_no: str) -> str:
    """측정일을 항상 YYYY-MM-DD로. 엑셀값 우선, 없으면 시료번호 YYMMDD."""
    s = _clean_cell(raw)
    if s:
        head = s.replace("/", "-").replace(".", "-")
        for fmt, n in (("%Y-%m-%d %H:%M:%S", 19), ("%Y-%m-%d", 10)):
            try:
                return datetime.strptime(head[:n], fmt).strftime("%Y-%m-%d")
            except ValueError:
                continue
        if len(head) >= 10 and head[4] == "-" and head[7] == "-":
            return head[:10]
    return (sample_to_datestr(sample_no) or "").strip()


def _resolve_air_meta(
    sample_no: str,
    excel_path: str,
    excel_meta: dict | None,
) -> dict:
    """
    excel_meta에 채취시간·날짜가 없으면 엑셀(대기측정기록부)에서 직접 채움.
    meta 없이 재전송해도 collected_at이 비지 않도록.
    """
    meta = dict(excel_meta or {})
    need_start = not _clean_cell(meta.get("채취시작"))
    need_end = not _clean_cell(meta.get("채취끝"))
    need_date = not _clean_cell(meta.get("날짜"))
    need_company = not _clean_cell(meta.get("업소명"))
    if not (need_start or need_end or need_date or need_company):
        return meta
    if not excel_path or not os.path.isfile(excel_path):
        return meta
    try:
        parsed = parse_measuring_record(excel_path, sample_no)
    except Exception as e:
        log_error("groupware._resolve_air_meta", e)
        return meta
    if need_start and parsed.get("채취시작"):
        meta["채취시작"] = parsed["채취시작"]
    if need_end and parsed.get("채취끝"):
        meta["채취끝"] = parsed["채취끝"]
    if need_date and parsed.get("날짜"):
        meta["날짜"] = parsed["날짜"]
    if need_company and parsed.get("업소명"):
        meta["업소명"] = parsed["업소명"]
    return meta


def _is_soft_facility_warning(w) -> bool:
    """
    시설명 매칭 실패 등 — 데이터는 이미 저장된 soft 경고.
    재시도해도 그룹웨어에 시설이 없으면 동일하므로 실패·재시도 대상에서 제외.
    """
    s = str(w or "")
    if "시설 연결 없이" in s or "시설명 매칭 실패" in s:
        return True
    sl = s.lower()
    return ("facility" in sl and ("match" in sl or "매칭" in s))


def _soft_facility_warnings(warnings: list) -> list:
    return [w for w in (warnings or []) if _is_soft_facility_warning(w)]


def _has_critical_field_warnings(warnings: list) -> bool:
    """채취시간·측정일 누락 등 재전송/수정이 필요한 hard 경고 (시설 soft 제외)."""
    for w in warnings or []:
        if _is_soft_facility_warning(w):
            continue
        s = str(w).lower()
        if any(
            k in s
            for k in (
                "collected_at",
                "measure_date",
                "채취시간",
                "측정일",
                "미전송",
            )
        ):
            return True
    return False


def _parse_api_body(resp: requests.Response) -> dict:
    """응답 JSON에서 status / warnings / verify_key / 매칭 업체 추출."""
    info = {
        "status": "",
        "warnings": [],
        "verify_key": "",
        "action": "",
        "message": "",
        "raw": "",
        "company_id": "",
        "matched_company_name": "",
        "company_matched_by": "",
        "facility_match": "",
    }
    try:
        data = resp.json()
    except Exception:
        info["raw"] = (resp.text or "")[:500]
        return info
    if not isinstance(data, dict):
        info["raw"] = str(data)[:500]
        return info
    info["status"] = str(data.get("status") or "")
    info["action"] = str(data.get("action") or "")
    info["verify_key"] = str(data.get("verify_key") or "")
    info["message"] = str(data.get("message") or data.get("detail") or "")
    cid = data.get("company_id")
    if cid not in (None, ""):
        info["company_id"] = str(cid)
    info["matched_company_name"] = str(data.get("matched_company_name") or "").strip()
    info["company_matched_by"] = str(data.get("company_matched_by") or "").strip()
    info["facility_match"] = str(data.get("facility_match") or "").strip()
    warns = data.get("warnings")
    if isinstance(warns, list):
        info["warnings"] = [str(w) for w in warns if w is not None]
    elif warns:
        info["warnings"] = [str(warns)]
    return info


def _format_api_matched_company(info: dict) -> str:
    """API 응답에서 실제 등록된 거래처(카드) 표시 문자열."""
    name = str(info.get("matched_company_name") or "").strip()
    if name:
        return name
    cid = str(info.get("company_id") or "").strip()
    if cid:
        return f"(company_id={cid})"
    return ""


def _print_api_feedback(sample_no: str, info: dict, label: str) -> None:
    status = info.get("status") or info.get("action") or ""
    vk = info.get("verify_key") or ""
    warns = info.get("warnings") or []
    matched = _format_api_matched_company(info)
    matched_by = str(info.get("company_matched_by") or "").strip()
    if status:
        extra = ""
        if vk:
            extra += f" verify_key={vk}"
        if matched:
            extra += f" api업체={matched!r}"
        if matched_by:
            extra += f" 매칭={matched_by}"
        print(f"   ← {label} status={status}{extra}")
    elif vk:
        print(f"   ← {label} verify_key={vk}")
    if warns:
        print(f"⚠ 그룹웨어 warnings ({sample_no}):")
        for w in warns:
            print(f"   · {w}")
        log_message(f"groupware.warnings {sample_no}: {warns}")


def _validate_payload_for_send(payload: dict) -> list[str]:
    """전송 전 필수 필드 검사. 빈 collected_at으로 덮어쓰지 않도록."""
    errs: list[str] = []
    if not _clean_cell(payload.get("sample_no")):
        errs.append("sample_no 없음")
    if not (
        _clean_cell(payload.get("company_name"))
        or _clean_cell(payload.get("biz_no"))
        or _clean_cell(payload.get("site_no"))
    ):
        errs.append("company_name/biz_no/site_no 없음")
    if not _clean_cell(payload.get("collected_at")):
        errs.append("collected_at(채취시간) 없음 — 엑셀 채취시작/끝을 확인하세요")
    if not _clean_cell(payload.get("measure_date")):
        errs.append("measure_date(측정일) 없음")
    if not (payload.get("measurements") or []):
        errs.append("measurements 비어 있음")
    return errs


def read_biz_no(excel_path: str) -> str:
    """사업자번호 — 성적서 엑셀에는 보통 없음. 보조 필드로만 비워 둠."""
    return ""


def _header_index_map(headers: list) -> dict[str, int]:
    m: dict[str, int] = {}
    for i, h in enumerate(headers or []):
        k = _norm_rg(h)
        if k and k not in m:
            m[k] = i
    return m


def _find_col_idx(hmap: dict[str, int], *candidates: str) -> int | None:
    for c in candidates:
        k = _norm_rg(c)
        if k in hmap:
            return hmap[k]
    for k, idx in hmap.items():
        for c in candidates:
            ck = _norm_rg(c)
            if ck and (ck in k or k in ck):
                return idx
    return None


def _cell_str(vals: list, idx: int | None) -> str:
    if idx is None or idx < 0 or idx >= len(vals):
        return ""
    v = vals[idx]
    return "" if v is None else str(v).strip()


def _conc_str(v) -> str:
    if v is None:
        return ""
    s = str(v).strip()
    if not s:
        return ""
    try:
        f = float(s.replace(",", ""))
        if f == int(f):
            return str(int(f))
        return str(f)
    except ValueError:
        return s


def _read_air_flow_rates(wb) -> tuple[str, str]:
    """
    대기측정기록부 — 배출가스유량 보정전(J14)·보정후(N14).
    excel_utils.parse_measuring_record 의 배출가스유량전/후 와 동일 셀.
    """
    ws = None
    for name in wb.sheetnames:
        if name.replace(" ", "") == "대기측정기록부":
            ws = wb[name]
            break
    if ws is None:
        return "", ""
    # J=10, N=14
    flow_pre = _conc_str(ws.cell(row=14, column=10).value)   # 보정전(미적용)
    flow_post = _conc_str(ws.cell(row=14, column=14).value)  # 보정후(적용1/적용2)
    return flow_pre, flow_post


def _norm_item_key(name: str) -> str:
    """항목명 매칭용 정규화 (공백 제거)."""
    return re.sub(r"\s+", "", str(name or "").strip())


def _read_record_item_units(wb) -> tuple[dict[str, str], dict[str, str]]:
    """
    대기측정기록부 측정결과 표 — B열 항목.
    - F열: 배출허용기준 단위 (standard_unit)
    - I열: 농도 단위 (unit)
    일반: '측정항목' 헤더 아래 구간.
    비산먼지: 상단 요약행(예: B19/F19/I19)도 포함.
    반환: (conc_unit_map, standard_unit_map) — 각각 {정규화항목명: 단위}
    """
    ws = None
    for name in wb.sheetnames:
        if name.replace(" ", "") == "대기측정기록부":
            ws = wb[name]
            break
    if ws is None:
        return {}, {}

    conc_out: dict[str, str] = {}
    std_out: dict[str, str] = {}

    def _add_row(r: int) -> bool:
        raw_item = ws.cell(row=r, column=2).value  # B
        if raw_item is None or not str(raw_item).strip():
            return False
        item_s = str(raw_item).strip()
        key = _norm_item_key(item_s)
        stop_keys = ("분석기간", "종합의견", "채취일시", "방지시설", "현장기상", "측정항목")
        if any(s in key for s in stop_keys):
            return False
        # F = 배출허용기준 단위, I = 농도 단위
        f_v = ws.cell(row=r, column=6).value  # F
        i_v = ws.cell(row=r, column=9).value  # I
        f_s = "" if f_v is None else str(f_v).strip()
        i_s = "" if i_v is None else str(i_v).strip()
        if f_s:
            std_out[key] = f_s
        if i_s:
            conc_out[key] = i_s
        return True

    # 비산먼지 양식: 측정항목 헤더 위쪽(대략 15~28행) B/F/I
    for r in range(15, 29):
        _add_row(r)

    header_rows: list[int] = []
    for r in range(20, 60):
        b = _norm_item_key(ws.cell(row=r, column=2).value)
        if "측정항목" in b:
            header_rows.append(r)
    if header_rows:
        start = header_rows[0] + 1
        end = header_rows[1] if len(header_rows) > 1 else min(start + 30, 80)
        for r in range(start, end):
            raw_item = ws.cell(row=r, column=2).value
            if raw_item is None or not str(raw_item).strip():
                continue
            key = _norm_item_key(str(raw_item).strip())
            if any(s in key for s in ("분석기간", "종합의견", "채취일시", "방지시설", "현장기상")):
                break
            _add_row(r)

    return conc_out, std_out


def _flow_for_air_ratio(air_ratio: str, flow_pre: str, flow_post: str) -> str:
    """공기비적용 값에 맞는 유량 선택. '적용' 포함(적용1·적용2) → 보정후, 그 외 → 보정전."""
    s = (air_ratio or "").strip()
    if "적용" in s and "미적용" not in s:
        return flow_post
    return flow_pre


def read_air_measurements(excel_path: str) -> list[dict]:
    """
    대기 — 입력(분석값)에서 항목·농도·기준치·공기비,
    대기측정기록부에서 농도단위(I)·기준단위(F)·유량(J14/N14).
    농도는 '농도'/'측정농도' 열. 배출허용기준(standard)은 '기준치'/'배출허용기준'/'배출허용기준농도' 열(없으면 H열).
    unit = 기록부 I열(농도 단위), standard_unit = 기록부 F열(배출허용기준 단위). B열 항목명 매칭.
    유량은 항목별 '공기비적용'(미적용/적용1/적용2)에 따라 J14 또는 N14.
    """
    out: list[dict] = []
    try:
        # read_only=False: 대기측정기록부 임의 셀(J14/N14) + 헤더 맵 안정 접근
        wb = load_workbook(excel_path, data_only=True, read_only=False)
    except Exception:
        return out

    try:
        flow_pre, flow_post = _read_air_flow_rates(wb)
        conc_unit_map, std_unit_map = _read_record_item_units(wb)

        ws = None
        for name in wb.sheetnames:
            if name.replace(" ", "") == "입력(분석값)":
                ws = wb[name]
                break
        if ws is None:
            for name in wb.sheetnames:
                if "입력" in name and "분석" in name:
                    ws = wb[name]
                    break
        if ws is None:
            return out

        header_map: dict[str, int] = {}
        for col in range(1, 40):
            v = ws.cell(row=1, column=col).value
            if v:
                header_map[str(v).strip()] = col

        def col_of(*names):
            for n in names:
                if n in header_map:
                    return header_map[n]
            return None

        c_stack = col_of("배출구", "굴뚝", "배출구번호", "시설", "시설명")
        c_unit = col_of("단위", "측정단위")
        c_conc = col_of("측정농도", "농도", "측정값", "분석결과")
        c_limit = col_of("기준치", "배출허용기준", "배출허용기준농도", "허용기준", "배출허용기준치", "permit_limit", "standard")
        c_air = col_of("공기비적용", "공기비")
        if not c_conc:
            c_conc = 5  # E열 관례
        if not c_limit:
            c_limit = 8  # H열 관례

        for r in range(2, 65):
            flag = ws.cell(row=r, column=1).value
            if flag is None or str(flag).strip() == "":
                continue
            item = ws.cell(row=r, column=2).value
            if not item:
                continue
            item_s = str(item).strip()
            if not item_s:
                continue

            stack = ""
            if c_stack:
                sv = ws.cell(row=r, column=c_stack).value
                stack = "" if sv is None else str(sv).strip()

            conc_val = ws.cell(row=r, column=c_conc).value if c_conc else None
            conc = _conc_str(conc_val)
            if not conc:
                continue

            item_key = _norm_item_key(item_s)
            unit = ""
            if c_unit:
                uv = ws.cell(row=r, column=c_unit).value
                unit = "" if uv is None else str(uv).strip()
            # 농도 단위: 입력(분석값) 단위 열 없으면 대기측정기록부 I열
            if not unit:
                unit = conc_unit_map.get(item_key, "")
            # 배출허용기준 단위: 대기측정기록부 F열
            standard_unit = std_unit_map.get(item_key, "")

            air_ratio = ""
            if c_air:
                av = ws.cell(row=r, column=c_air).value
                air_ratio = "" if av is None else str(av).strip()

            limit = ""
            if c_limit:
                lv = ws.cell(row=r, column=c_limit).value
                limit = _conc_str(lv)

            out.append({
                "stack": stack,
                "item_name": item_s,
                "concentration": conc,
                "standard": limit,
                "standard_unit": standard_unit,
                "unit": unit,
                "flow_rate": _flow_for_air_ratio(air_ratio, flow_pre, flow_post),
            })
    except Exception as e:
        log_error("groupware.read_air_measurements", e)
    finally:
        try:
            wb.close()
        except Exception:
            pass
    return out


def read_water_measurements(tab4_meta: dict) -> list[dict]:
    """수질 — 탭4 매크로 headers/rows에서 항목·농도 추출."""
    headers = tab4_meta.get("headers") or []
    rows = tab4_meta.get("rows") or []
    hmap = _header_index_map(headers)

    i_item = _find_col_idx(hmap, "측정항목", "항목", "분석항목")
    i_stack = _find_col_idx(hmap, "배출구", "굴뚝", "시설", "시설명", "측정지점")
    i_conc = _find_col_idx(hmap, "측정농도", "농도", "측정값", "분석결과", "결과")
    i_unit = _find_col_idx(hmap, "단위", "측정단위")

    out: list[dict] = []
    for row_vals in rows:
        item = _cell_str(row_vals, i_item)
        if not item and row_vals:
            item = _norm_rg(row_vals[0])
            if len(row_vals) > 1 and not item:
                item = _norm_rg(row_vals[1])
        if not item:
            continue

        conc = _conc_str(_cell_str(row_vals, i_conc) if i_conc is not None else "")
        if not conc and i_conc is None and len(row_vals) > 4:
            conc = _conc_str(row_vals[4])

        if not conc:
            continue

        out.append({
            "stack": _cell_str(row_vals, i_stack),
            "item_name": item,
            "concentration": conc,
            "unit": _cell_str(row_vals, i_unit),
        })
    return out


def _first_stack(measurements: list[dict]) -> str:
    for m in measurements:
        s = (m.get("stack") or "").strip()
        if s:
            return s
    return ""


def _merge_input_keys(excel_path: str, excel_meta: dict | None) -> dict:
    """
    입력 시트 사업장명(H7)·사업장관리번호(H6)·측정인 시설명(E4) 우선.
    excel_meta 업소명은 H7이 비었을 때만 보조.
    """
    keys = read_input_sheet_keys(excel_path)
    if excel_meta and not keys["company_name"]:
        keys["company_name"] = _clean_cell(excel_meta.get("업소명"))
    if not keys.get("biz_no"):
        keys["biz_no"] = read_biz_no(excel_path)
    return keys


def build_payload_air(
    sample_no: str,
    excel_path: str,
    excel_meta: dict | None = None,
    tab4_meta: dict | None = None,
) -> dict:
    """
    대기 성적서 → 그룹웨어 ReportDataIn.
    collected_at / measure_date 는 excel_meta가 비어도 엑셀에서 채움.

    거래처·시설 매칭 (수신: ECO_INPUT_MATCHING.md):
      site_no ← 입력!H6 사업장관리번호 — 있으면 거래처 매칭 최우선
      company_name ← 입력!H7 업체명 = 마스터 G 전체 (짧은 이름·의뢰기관 C 금지)
      facility_name ← 입력!E4 측정인 시설명 = 마스터 K
      sems_stack ← 입력!E2 (facility_name 실패·공란 시 폴백)
      biz_no ← 있으면 보조만 (매칭 의존 금지)
    """
    meta = _resolve_air_meta(sample_no, excel_path, excel_meta)
    input_keys = _merge_input_keys(excel_path, meta)
    fac = facility_match_keys(input_keys)

    measurements = read_air_measurements(excel_path)

    collect_start = format_time(meta.get("채취시작", "")) if meta.get("채취시작") else ""
    collect_end = format_time(meta.get("채취끝", "")) if meta.get("채취끝") else ""
    collected_at = format_collected_at(collect_start, collect_end)
    measure_date = _normalize_measure_date(meta.get("날짜"), sample_no)

    facility_name = fac.get("facility_name") or ""

    return {
        "media": "air",
        "sample_no": sample_no,
        "collected_at": collected_at,
        "collect_start": collect_start,
        "collect_end": collect_end,
        "measure_date": measure_date,
        **fac,
        "facility_name": facility_name,
        "measurements": measurements,
        "source_excel": os.path.abspath(excel_path) if excel_path else "",
    }


# ---------------------------------------------------------------------------
# 수질 — 그룹웨어 연동 미적용 (대기 E2/E4 시설매칭 규칙 해당 없음)
# 프롬프트 3 수질 경로는 eco_input._main_water 에서 호출 주석 처리됨.
# ---------------------------------------------------------------------------
def build_payload_water(
    sample_no: str,
    excel_path: str,
    tab4_meta: dict,
    company_name: str = "",
) -> dict:
    """미사용 — 수질은 그룹웨어 연동 대상 아님. 호출 시 명시적 오류."""
    raise NotImplementedError(
        "수질 성적서는 그룹웨어(프롬프트3) 연동 대상이 아닙니다. 대기만 사용하세요."
    )


# ---------------------------------------------------------------------------
# API 전송
# ---------------------------------------------------------------------------
def _request_with_retry(method: str, url: str, **kwargs) -> requests.Response | None:
    retries = max(1, int(GROUPWARE_MAX_RETRIES))
    timeout = float(kwargs.pop("timeout", GROUPWARE_TIMEOUT_SEC))
    last_err = None

    for attempt in range(1, retries + 1):
        try:
            fn = requests.post if method.upper() == "POST" else requests.get
            resp = fn(url, timeout=timeout, **kwargs)
            if resp.status_code >= 500 and attempt < retries:
                time.sleep(1.0 * attempt)
                continue
            return resp
        except Exception as e:
            last_err = e
            log_error(f"groupware.{method} {url} (attempt {attempt})", e)
            if attempt < retries:
                time.sleep(1.0 * attempt)

    if last_err:
        raise last_err
    return None


def post_report_data(payload: dict) -> tuple[bool, str, dict]:
    """POST /api/external/report_data → (ok, error, api_info)."""
    url = _url(GROUPWARE_REPORT_DATA_PATH)
    empty_info: dict = {"status": "", "warnings": [], "verify_key": ""}
    try:
        resp = _request_with_retry(
            "POST",
            url,
            headers={**_auth_headers(), "Content-Type": "application/json"},
            json=payload,
        )
        if resp is None:
            return False, "no response", empty_info
        info = _parse_api_body(resp)
        if resp.status_code >= 400:
            msg = info.get("message") or info.get("raw") or resp.text[:300]
            print(f"⚠ 그룹웨어 report_data 실패 ({resp.status_code}): {msg}")
            return False, msg, info
        print(f"✅ 그룹웨어 데이터 전송: {payload.get('sample_no')}")
        _print_api_feedback(payload.get("sample_no", ""), info, "report_data")
        return True, "", info
    except Exception as e:
        print(f"⚠ 그룹웨어 report_data 오류 ({payload.get('sample_no')}): {e}")
        log_error("groupware.post_report_data", e)
        return False, str(e), empty_info


def post_report_sync(payload: dict, pdf_path: str) -> tuple[bool, str, dict]:
    """POST /api/reports/sync — JSON + PDF 한 번에 → (ok, error, api_info)."""
    empty_info: dict = {"status": "", "warnings": [], "verify_key": ""}
    if not pdf_path or not os.path.isfile(pdf_path):
        return False, "pdf missing", empty_info

    url = _url(GROUPWARE_REPORT_SYNC_PATH)
    data_json = json.dumps(payload, ensure_ascii=False)
    try:
        with open(pdf_path, "rb") as f:
            resp = _request_with_retry(
                "POST",
                url,
                headers=_auth_headers(),
                data={"data": data_json},
                files={
                    "file": (
                        os.path.basename(pdf_path),
                        f,
                        "application/pdf",
                    )
                },
                timeout=max(float(GROUPWARE_TIMEOUT_SEC), 120.0),
            )
        if resp is None:
            return False, "no response", empty_info
        info = _parse_api_body(resp)
        if resp.status_code >= 400:
            msg = info.get("message") or info.get("raw") or resp.text[:300]
            print(f"⚠ 그룹웨어 reports/sync 실패 ({resp.status_code}): {msg}")
            return False, msg, info
        print(f"✅ 그룹웨어 sync(PDF+데이터): {payload.get('sample_no')}")
        _print_api_feedback(payload.get("sample_no", ""), info, "reports/sync")
        return True, "", info
    except Exception as e:
        print(f"⚠ 그룹웨어 reports/sync 오류 ({payload.get('sample_no')}): {e}")
        log_error("groupware.post_report_sync", e)
        return False, str(e), empty_info


def post_report_pdf(payload: dict, pdf_path: str) -> tuple[bool, str, dict]:
    """POST /api/external/report_pdf — PDF만 → (ok, error, api_info)."""
    empty_info: dict = {"status": "", "warnings": [], "verify_key": ""}
    if not pdf_path or not os.path.isfile(pdf_path):
        return False, "pdf missing", empty_info

    url = _url(GROUPWARE_REPORT_PDF_PATH)
    # ECO_INPUT_MATCHING.md: report_pdf Form 에도 site_no·facility_name 권장
    form = {
        "sample_no": payload.get("sample_no", ""),
        "media": payload.get("media", ""),
        "company_name": payload.get("company_name", ""),
        "site_no": payload.get("site_no", ""),
        "biz_no": payload.get("biz_no", ""),
        "facility_name": payload.get("facility_name", ""),
        "sems_stack": payload.get("sems_stack") or payload.get("sems_serial", ""),
        "sems_serial": payload.get("sems_serial", ""),
        "facility_alt": payload.get("facility_alt", ""),
        "collected_at": payload.get("collected_at", ""),
        "measure_date": payload.get("measure_date", ""),
    }
    try:
        with open(pdf_path, "rb") as f:
            resp = _request_with_retry(
                "POST",
                url,
                headers=_auth_headers(),
                data=form,
                files={
                    "file": (
                        os.path.basename(pdf_path),
                        f,
                        "application/pdf",
                    )
                },
                timeout=max(float(GROUPWARE_TIMEOUT_SEC), 120.0),
            )
        if resp is None:
            return False, "no response", empty_info
        info = _parse_api_body(resp)
        if resp.status_code >= 400:
            msg = info.get("message") or info.get("raw") or resp.text[:300]
            print(f"⚠ 그룹웨어 report_pdf 실패 ({resp.status_code}): {msg}")
            return False, msg, info
        print(f"✅ 그룹웨어 PDF 전송: {payload.get('sample_no')}")
        _print_api_feedback(payload.get("sample_no", ""), info, "report_pdf")
        return True, "", info
    except Exception as e:
        print(f"⚠ 그룹웨어 report_pdf 오류 ({payload.get('sample_no')}): {e}")
        log_error("groupware.post_report_pdf", e)
        return False, str(e), empty_info


def sync_to_groupware(payload: dict, pdf_path: str | None = None) -> dict:
    """
    그룹웨어 전송 (실패해도 예외 없음).
    반환: {data_ok, pdf_ok, pdf_attempted, error, warnings, verify_key, status}
    """
    result = {
        "data_ok": False,
        "pdf_ok": False,
        "pdf_attempted": bool(pdf_path),
        "error": "",
        "warnings": [],
        "verify_key": "",
        "status": "",
        "company_id": "",
        "matched_company_name": "",
        "company_matched_by": "",
        "facility_match": "",
    }

    # 빈 collected_at 재전송 금지 (기존 DB 값 지우기 방지 + 최초 누락 방지)
    val_errs = _validate_payload_for_send(payload)
    if val_errs:
        msg = "; ".join(val_errs)
        print(f"❌ 그룹웨어 전송 중단 ({payload.get('sample_no')}): {msg}")
        result["error"] = msg
        result["warnings"] = val_errs
        return result

    print(
        f"▶ 그룹웨어 payload: sample_no={payload.get('sample_no')} "
        f"collected_at={payload.get('collected_at')!r} "
        f"measure_date={payload.get('measure_date')!r} "
        f"facility_name={payload.get('facility_name')!r} "
        f"company={payload.get('company_name')!r} "
        f"site_no={payload.get('site_no')!r}"
    )
    log_message(
        f"groupware.send {payload.get('sample_no')} "
        f"collected_at={payload.get('collected_at')} "
        f"measure_date={payload.get('measure_date')} "
        f"facility_name={payload.get('facility_name')}"
    )

    def _merge_info(info: dict) -> None:
        result["warnings"] = list(info.get("warnings") or [])
        result["verify_key"] = info.get("verify_key") or result["verify_key"]
        result["status"] = info.get("status") or info.get("action") or result["status"]
        if info.get("company_id"):
            result["company_id"] = str(info.get("company_id") or "")
        if info.get("matched_company_name"):
            result["matched_company_name"] = str(info.get("matched_company_name") or "")
        if info.get("company_matched_by"):
            result["company_matched_by"] = str(info.get("company_matched_by") or "")
        if info.get("facility_match"):
            result["facility_match"] = str(info.get("facility_match") or "")

    if pdf_path and os.path.isfile(pdf_path):
        ok, err, info = post_report_sync(payload, pdf_path)
        _merge_info(info)
        if ok:
            result["data_ok"] = True
            result["pdf_ok"] = True
            if _has_critical_field_warnings(result["warnings"]):
                result["error"] = "API warnings: " + "; ".join(result["warnings"])
            return result
        result["error"] = err

    ok, err, info = post_report_data(payload)
    _merge_info(info)
    result["data_ok"] = ok
    if not ok:
        result["error"] = err
        return result

    if _has_critical_field_warnings(result["warnings"]):
        result["error"] = "API warnings: " + "; ".join(result["warnings"])

    if pdf_path and os.path.isfile(pdf_path):
        pok, perr, pinfo = post_report_pdf(payload, pdf_path)
        # PDF 응답 warnings는 합침
        extra = list(pinfo.get("warnings") or [])
        if extra:
            result["warnings"] = list(result["warnings"]) + extra
        if pinfo.get("verify_key"):
            result["verify_key"] = pinfo["verify_key"]
        result["pdf_ok"] = pok
        if not pok:
            result["error"] = perr

    return result


def payload_to_excel_rows(payload: dict, sync_result: dict) -> list[dict]:
    """로컬 정리 엑셀용 flat 행 (항목별 1행)."""
    warns = sync_result.get("warnings") or []
    soft = _soft_facility_warnings(warns)
    hard = _has_critical_field_warnings(warns)
    data_ok = bool(sync_result.get("data_ok"))
    pdf_ok = bool(sync_result.get("pdf_ok"))
    pdf_att = bool(sync_result.get("pdf_attempted"))
    need_resend = (not data_ok) or (pdf_att and not pdf_ok) or hard or bool(soft)
    note = sync_result.get("error", "") or (
        "; ".join(str(w) for w in soft) if soft else ""
    )
    api_company = _format_api_matched_company(sync_result)
    api_match = str(sync_result.get("company_matched_by") or "").strip()
    base = {
        "시료번호": payload.get("sample_no", ""),
        "매체": payload.get("media", ""),
        "업소명": payload.get("company_name", ""),
        "업소명_정규화": payload.get("company_name_norm", ""),
        "API등록업체": api_company,
        "API매칭방식": api_match,
        "사업장관리번호": payload.get("site_no", ""),
        "SEMS일련번호_E2": payload.get("sems_serial", ""),
        "시설매칭보조_E4": payload.get("facility_alt", ""),
        "채취시간": payload.get("collected_at", ""),
        "측정일": payload.get("measure_date", ""),
        "시설명": payload.get("facility_name", ""),
        "데이터전송": "OK" if data_ok else "FAIL",
        "PDF전송": (
            "OK" if pdf_ok
            else ("N/A" if not pdf_att else "FAIL")
        ),
        "비고": note,
        "원본엑셀": payload.get("source_excel", ""),
        "재전송상태": RESEND_STATUS_PENDING if need_resend else "",
        "재전송시각": "",
    }
    measurements = payload.get("measurements") or []
    if not measurements:
        return [dict(base, 굴뚝="", 항목="", 농도="", 기준치="", 기준단위="", 단위="", 유량="")]

    rows = []
    for m in measurements:
        rows.append(dict(
            base,
            굴뚝=m.get("stack", ""),
            항목=m.get("item_name", ""),
            농도=m.get("concentration", ""),
            기준치=m.get("standard", "") or m.get("limit", ""),
            기준단위=m.get("standard_unit", ""),
            단위=m.get("unit", ""),
            유량=m.get("flow_rate", ""),
        ))
    return rows


def sync_tab4_complete(
    gw_log: GroupwareRunLog | None,
    *,
    media: str,
    sample_no: str,
    excel_path: str,
    tab4_meta: dict | None = None,
    excel_meta: dict | None = None,
    pdf_path: str | None = None,
    company_name: str = "",
    _log_to_gw_log: bool = True,
) -> dict:
    """
    탭4 입력완료 직후 호출. payload 조립 → API 전송 → gw_log 누적.
    **대기(air) 전용** — 수질은 연동하지 않음.
    _log_to_gw_log=False 로 호출하면 gw_log 누적을 건너뜀 (재시도 중간 호출용).
    """
    if media == "water":
        print("⚠ 그룹웨어 연동: 수질은 미지원 (스킵)")
        return {"data_ok": False, "pdf_ok": False, "pdf_attempted": False, "error": "water skipped"}

    payload = build_payload_air(
        sample_no, excel_path, excel_meta=excel_meta, tab4_meta=tab4_meta
    )

    print(
        f"▶ 그룹웨어 전송 준비: {sample_no} "
        f"(항목 {len(payload.get('measurements', []))}개, "
        f"채취={payload.get('collected_at')!r}, "
        f"측정일={payload.get('measure_date')!r}, "
        f"사업장={payload.get('company_name')!r}, "
        f"site_no={payload.get('site_no')!r}, "
        f"시설명={payload.get('facility_name')!r})"
    )
    sync_result = sync_to_groupware(payload, pdf_path=pdf_path)

    if gw_log is not None:
        if _log_to_gw_log:
            gw_log.add_rows(payload_to_excel_rows(payload, sync_result))
        # 결과는 항상 기록 (재시도 시 마지막 결과로 덮어씀)
        gw_log.record_result(sample_no, sync_result, payload)

    # payload를 결과에 붙여 재시도 시 재사용 가능하게
    sync_result["_payload"] = payload
    return sync_result


def export_groupware_summary(
    gw_log: GroupwareRunLog | None,
    out_dir: str | None = None,
) -> str | None:
    """
    전송 기록을 로컬 엑셀로 저장 (오토핏).
    경로: {EXCEL_OUTPUT_DIR}/{YYYY}년/{M}월/그룹웨어전송_YYYYMMDD_HHMMSS.xlsx
    """
    if gw_log is None or not gw_log.rows:
        return None

    now = datetime.now()
    base_dir = out_dir or GROUPWARE_EXCEL_OUTPUT_DIR
    # 6.그룹웨어전송\2026년\7월\...
    out_dir = os.path.join(base_dir, f"{now.year}년", f"{now.month}월")
    os.makedirs(out_dir, exist_ok=True)

    stamp = now.strftime("%Y%m%d_%H%M%S")
    out_path = os.path.join(out_dir, f"그룹웨어전송_{stamp}.xlsx")

    headers = [
        "시료번호", "매체", "업소명", "업소명_정규화",
        "API등록업체", "API매칭방식",
        "사업장관리번호",
        "SEMS일련번호_E2", "시설매칭보조_E4",
        "채취시간", "측정일", "시설명",
        "굴뚝", "항목", "농도", "기준치", "기준단위", "단위", "유량", "데이터전송", "PDF전송", "비고", "원본엑셀",
        "재전송상태", "재전송시각",
    ]

    wb = Workbook()
    ws = wb.active
    ws.title = "전송기록"
    ws.append(headers)
    for row in gw_log.rows:
        ws.append([row.get(h, "") for h in headers])

    last_col = chr(ord("A") + len(headers) - 1)
    autofit_columns(ws, f"A:{last_col}")
    wb.save(out_path)
    wb.close()

    print(f"✅ 그룹웨어 전송 정리 엑셀: {out_path}")
    return out_path


# ---------------------------------------------------------------------------
# 기간별 재전송 (나중에 모아서 처리)
# ---------------------------------------------------------------------------
def _classify_sample_resend_need(
    s: dict,
    *,
    include_facility_soft: bool,
    include_pdf_missing: bool = True,
) -> str:
    """
    재전송 사유 문자열. 불필요하면 빈 문자열.
    스킵은 제외. PDF 없음/실패는 재전송상태=완료여도 포함.
    """
    st = (s.get("resend_status") or "").strip()
    if st == RESEND_STATUS_SKIP:
        return ""

    pdf_ok = bool(s.get("pdf_ok"))
    pdf_attempted = bool(s.get("pdf_attempted"))
    pdf_missing = (not pdf_ok) and (not pdf_attempted)
    pdf_fail = pdf_attempted and (not pdf_ok)

    # 재전송 완료면 예전 로그의 데이터전송=FAIL 은 무시.
    # PDF 없음/실패만 다시 대상으로 남긴다.
    if st == RESEND_STATUS_DONE:
        if pdf_fail:
            return "PDF FAIL"
        if include_pdf_missing and pdf_missing:
            return "PDF 없음"
        return ""

    if not s.get("data_ok"):
        return "데이터 FAIL"
    if pdf_fail:
        return "PDF FAIL"
    if include_pdf_missing and pdf_missing:
        return "PDF 없음"

    note = s.get("note") or ""
    soft_only = bool(_soft_facility_warnings([note])) and s.get("data_ok") and (
        s.get("pdf_ok") or not s.get("pdf_attempted")
    )
    if note and _has_critical_field_warnings([note]):
        return "필드 경고"
    if include_facility_soft and soft_only:
        return "시설 soft"
    if st == RESEND_STATUS_PENDING:
        return "대기"
    return ""


def _parse_log_stamp_from_name(path: str) -> datetime | None:
    """그룹웨어전송_YYYYMMDD_HHMMSS.xlsx → datetime."""
    base = os.path.basename(path or "")
    m = re.search(r"그룹웨어전송_(\d{8})_(\d{6})", base)
    if not m:
        return None
    try:
        return datetime.strptime(m.group(1) + m.group(2), "%Y%m%d%H%M%S")
    except ValueError:
        return None


def list_groupware_log_files(
    date_from: datetime | str | None = None,
    date_to: datetime | str | None = None,
    base_dir: str | None = None,
) -> list[str]:
    """기간 내 그룹웨어전송_*.xlsx 경로 목록 (오래된 순)."""
    root = base_dir or GROUPWARE_EXCEL_OUTPUT_DIR
    if not root or not os.path.isdir(root):
        return []

    def _as_dt(v, end=False):
        if v is None:
            return None
        if isinstance(v, datetime):
            return v
        s = str(v).strip().replace("/", "-").replace(".", "-")
        for fmt, n in (("%Y-%m-%d %H:%M:%S", 19), ("%Y-%m-%d", 10)):
            try:
                d = datetime.strptime(s[:n], fmt)
                if end and fmt == "%Y-%m-%d":
                    return d.replace(hour=23, minute=59, second=59)
                return d
            except ValueError:
                continue
        return None

    d0 = _as_dt(date_from, end=False)
    d1 = _as_dt(date_to, end=True)

    out = []
    for dirpath, _dirs, files in os.walk(root):
        for fn in files:
            if not fn.startswith("그룹웨어전송_") or not fn.lower().endswith(".xlsx"):
                continue
            if fn.startswith("~$"):
                continue
            path = os.path.join(dirpath, fn)
            stamp = _parse_log_stamp_from_name(path)
            if stamp is None:
                # 파일 mtime fallback
                try:
                    stamp = datetime.fromtimestamp(os.path.getmtime(path))
                except OSError:
                    continue
            if d0 and stamp < d0:
                continue
            if d1 and stamp > d1:
                continue
            out.append((stamp, path))
    out.sort(key=lambda x: x[0])
    return [p for _, p in out]


def _iter_summary_excel_samples(excel_path: str) -> list[dict]:
    """
    그룹웨어전송_*.xlsx 에서 시료별 1건씩 요약.
    """
    wb = load_workbook(excel_path, data_only=True, read_only=True)
    ws = wb.active
    rows_iter = ws.iter_rows(values_only=True)
    try:
        headers = [str(c or "").strip() for c in next(rows_iter)]
    except StopIteration:
        wb.close()
        return []
    idx = {h: i for i, h in enumerate(headers)}

    def cell(row, key, default=""):
        i = idx.get(key)
        if i is None or i >= len(row):
            return default
        v = row[i]
        return "" if v is None else str(v).strip()

    by_sample: dict[str, dict] = {}
    for row in rows_iter:
        sno = cell(row, "시료번호")
        if not sno:
            continue
        data = cell(row, "데이터전송").upper()
        pdf = cell(row, "PDF전송").upper()
        note = cell(row, "비고")
        src = cell(row, "원본엑셀")
        fac = cell(row, "시설명")
        company = cell(row, "업소명")
        api_company = cell(row, "API등록업체")
        api_match = cell(row, "API매칭방식")
        rst = cell(row, "재전송상태")
        rtm = cell(row, "재전송시각")
        if sno not in by_sample:
            by_sample[sno] = {
                "sample_no": sno,
                "data_ok": data == "OK",
                "pdf_ok": pdf == "OK",
                "pdf_attempted": pdf not in ("", "N/A", "NA"),
                "note": note,
                "source_excel": src,
                "facility_name": fac,
                "company_name": company,
                "api_matched_company": api_company,
                "company_matched_by": api_match,
                "resend_status": rst,
                "resend_time": rtm,
                "log_path": excel_path,
            }
        else:
            cur = by_sample[sno]
            if data != "OK":
                cur["data_ok"] = False
            if pdf == "OK":
                cur["pdf_ok"] = True
                cur["pdf_attempted"] = True
            if pdf == "FAIL":
                cur["pdf_ok"] = False
                cur["pdf_attempted"] = True
            if note and not cur["note"]:
                cur["note"] = note
            if src and not cur["source_excel"]:
                cur["source_excel"] = src
            if fac and not cur.get("facility_name"):
                cur["facility_name"] = fac
            if company and not cur.get("company_name"):
                cur["company_name"] = company
            if api_company:
                cur["api_matched_company"] = api_company
            if api_match:
                cur["company_matched_by"] = api_match
            if rst and not cur["resend_status"]:
                cur["resend_status"] = rst
            if rtm and not cur["resend_time"]:
                cur["resend_time"] = rtm
    wb.close()
    return list(by_sample.values())


def collect_pending_resends(
    date_from: datetime | str | None = None,
    date_to: datetime | str | None = None,
    *,
    include_facility_soft: bool = True,
    include_pdf_missing: bool = True,
    base_dir: str | None = None,
) -> list[dict]:
    """
    기간 내 전송 로그를 모아 재전송 대상(아직 완료/스킵 아닌 건)을 반환.
    PDF 없음(N/A)·PDF FAIL 은 재전송상태=완료여도 포함.
    동일 시료가 여러 로그에 있으면 **최신 로그** 기준.
    """
    files = list_groupware_log_files(date_from, date_to, base_dir=base_dir)
    latest: dict[str, dict] = {}
    for path in files:
        stamp = _parse_log_stamp_from_name(path)
        for s in _iter_summary_excel_samples(path):
            sno = s["sample_no"]
            s = dict(s)
            s["log_path"] = path
            s["log_stamp"] = stamp.strftime("%Y-%m-%d %H:%M:%S") if stamp else ""
            reason = _classify_sample_resend_need(
                s,
                include_facility_soft=include_facility_soft,
                include_pdf_missing=include_pdf_missing,
            )
            if not reason:
                # 시간순 처리: 이후(또는 현재) 로그가 성공·완료면 이전 대기 건 제외
                latest.pop(sno, None)
                continue
            s["reason"] = reason
            s["_stamp"] = stamp or datetime.min
            latest[sno] = s

    out = []
    for s in latest.values():
        s.pop("_stamp", None)
        out.append(s)
    out.sort(key=lambda x: (x.get("log_stamp") or "", x.get("sample_no") or ""))
    print(
        f"▶ 재전송 대상 조회: {len(out)}건 "
        f"(로그파일 {len(files)}개, "
        f"soft={'포함' if include_facility_soft else '제외'}, "
        f"PDF없음={'포함' if include_pdf_missing else '제외'})"
    )
    return out


def mark_resend_status_in_log(
    log_path: str,
    sample_nos: list[str] | set[str],
    status: str = RESEND_STATUS_DONE,
    *,
    note_append: str = "",
    pdf_status: str | None = None,
    data_status: str | None = None,
) -> int:
    """
    전송 로그 엑셀에서 해당 시료 행의 재전송상태/시각을 갱신.
    pdf_status / data_status 가 있으면 PDF전송·데이터전송 열도 같이 갱신.
    열이 없으면 추가. 갱신한 행 수 반환.
    """
    want = {str(s).strip() for s in sample_nos if s}
    if not want or not log_path or not os.path.isfile(log_path):
        return 0

    wb = load_workbook(log_path)
    ws = wb.active
    headers = [str(c.value or "").strip() for c in ws[1]]
    header_map = {h: i + 1 for i, h in enumerate(headers)}  # 1-based

    def _ensure_col(name: str) -> int:
        if name in header_map:
            return header_map[name]
        col = len(headers) + 1
        ws.cell(1, col, name)
        headers.append(name)
        header_map[name] = col
        return col

    c_sample = header_map.get("시료번호")
    if not c_sample:
        wb.close()
        return 0
    c_status = _ensure_col("재전송상태")
    c_time = _ensure_col("재전송시각")
    c_note = header_map.get("비고")
    c_pdf = header_map.get("PDF전송") if pdf_status else None
    c_data = header_map.get("데이터전송") if data_status else None
    # 원본에 대기 표시가 없었을 수도 있음 — 완료만 기록

    now_s = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    n = 0
    for r in range(2, ws.max_row + 1):
        sno = str(ws.cell(r, c_sample).value or "").strip()
        if sno not in want:
            continue
        ws.cell(r, c_status, status)
        ws.cell(r, c_time, now_s)
        if pdf_status and c_pdf:
            ws.cell(r, c_pdf, pdf_status)
        if data_status and c_data:
            ws.cell(r, c_data, data_status)
        if note_append and c_note:
            old = str(ws.cell(r, c_note).value or "").strip()
            if note_append not in old:
                ws.cell(r, c_note, (old + " | " if old else "") + note_append)
        n += 1

    if n:
        try:
            last_col = chr(ord("A") + max(len(headers) - 1, 0))
            autofit_columns(ws, f"A:{last_col}")
        except Exception:
            pass
        wb.save(log_path)
    wb.close()
    return n


def map_report_excels_by_sample(paths: list[str]) -> dict[str, str]:
    """성적서 파일 경로 목록 → {시료번호: 경로}."""
    from data_utils import extract_sample_from_name

    out: dict[str, str] = {}
    for p in paths or []:
        p = str(p or "").strip()
        if not p or not os.path.isfile(p):
            continue
        sno = extract_sample_from_name(p)
        if sno:
            out[sno] = p
    return out


def _make_groupware_pdf(src: str, sno: str) -> tuple[str | None, str]:
    """
    성적서에서 그룹웨어용 PDF 생성.
    반환: (pdf_path 또는 None, error). 성공 시 error는 빈 문자열.
    """
    from eco_input import PDF_TMP_DIR, cleanup_tmp_pdfs, make_tab4_pdfs

    try:
        print(f"  → {sno}: PDF 생성 중... ({os.path.basename(src)})")
        pdfs = make_tab4_pdfs(src, sno, copy_to_nas=False)
        pdf_path = pdfs.get("pdf_groupware") or pdfs.get("pdf_record")
        if pdf_path and os.path.isfile(pdf_path):
            print(f"     PDF OK: {os.path.basename(pdf_path)}")
            return pdf_path, ""
        print("     ❌ 그룹웨어 PDF 경로 없음")
        return None, "그룹웨어 PDF 생성 결과 없음"
    except Exception as e:
        print(f"  ❌ {sno}: PDF 생성 실패 — {e}")
        log_error("groupware.make_groupware_pdf", e)
        try:
            cleanup_tmp_pdfs(PDF_TMP_DIR, sno)
        except Exception:
            pass
        return None, f"PDF 생성 실패: {e}"


def _sync_excel_with_pdf(sno: str, src: str, gw_log: GroupwareRunLog) -> dict:
    """
    성적서 1건을 eco_input과 동일하게 데이터+PDF 전송.
    gw_log._results[sno] 기준으로 판정 dict 반환.
    """
    from eco_input import (
        PDF_TMP_DIR,
        _try_groupware_tab4_sync,
        cleanup_tmp_pdfs,
    )

    pdf_path, pdf_err = _make_groupware_pdf(src, sno)
    if pdf_err or not pdf_path:
        return {
            "sample_no": sno, "ok": False, "soft_only": False,
            "error": pdf_err or "그룹웨어 PDF 없음",
            "warnings": [], "verify_key": "",
            "source_excel": src, "data_ok": False, "pdf_ok": False,
        }

    try:
        _try_groupware_tab4_sync(
            gw_log,
            media="air",
            sample_no=sno,
            excel_path=src,
            tab4_meta=None,
            excel_meta=None,
            pdf_path=pdf_path,
        )
    except Exception as e:
        print(f"  ❌ {sno}: 전송 오류 — {e}")
        log_error("groupware.sync_excel_with_pdf", e)
        return {
            "sample_no": sno, "ok": False, "soft_only": False,
            "error": str(e), "warnings": [], "verify_key": "",
            "source_excel": src, "data_ok": False, "pdf_ok": False,
        }
    finally:
        try:
            cleanup_tmp_pdfs(PDF_TMP_DIR, sno)
        except Exception:
            pass

    r = (gw_log._results or {}).get(sno) or {}
    warns = r.get("warnings") or []
    soft = _soft_facility_warnings(warns)
    hard = _has_critical_field_warnings(warns)
    data_ok = bool(r.get("data_ok"))
    pdf_ok = bool(r.get("pdf_ok"))
    pipeline_ok = data_ok and pdf_ok and not hard
    soft_only = pipeline_ok and bool(soft)
    ok = pipeline_ok and not soft
    return {
        "sample_no": sno,
        "ok": ok,
        "soft_only": soft_only,
        "error": r.get("error", ""),
        "warnings": warns,
        "verify_key": r.get("verify_key", ""),
        "source_excel": src,
        "data_ok": data_ok,
        "pdf_ok": pdf_ok,
    }


def send_from_report_excels(
    paths: list[str],
    *,
    write_summary: bool = True,
    progress_cb=None,
) -> list[dict]:
    """
    성적서 엑셀만으로 그룹웨어 전송 (전송로그 조회 불필요).
    eco_input 탭4 완료 후와 동일: make_tab4_pdfs → _try_groupware_tab4_sync
    (데이터 + 그룹웨어 PDF, 재시도 포함).
    write_summary=True 이면 6.그룹웨어전송 에 전송기록 엑셀도 남김.
    progress_cb: 선택. (cur, total, sample_no, status) 호출.
    """
    path_map = map_report_excels_by_sample(paths)
    if not path_map:
        print("▶ 직접 전송 대상 없음 (시료번호 추출 실패 또는 파일 없음)")
        return []

    items = sorted(path_map.items())
    total = len(items)
    print(f"▶ 그룹웨어 성적서 직접 전송 시작 ({total}건) — 데이터+PDF (eco_input과 동일)")
    gw_log = GroupwareRunLog()
    results = []
    for i, (sno, src) in enumerate(items, start=1):
        if progress_cb:
            try:
                progress_cb(i, total, sno, "전송 중")
            except Exception:
                pass
        results.append(_sync_excel_with_pdf(sno, src, gw_log))
        time.sleep(1.5)

    if write_summary and gw_log.rows:
        export_groupware_summary(gw_log)

    ok_n = sum(1 for r in results if r.get("ok"))
    soft_n = sum(1 for r in results if r.get("soft_only"))
    ng = sum(1 for r in results if not r.get("ok") and not r.get("soft_only"))
    print(
        f"✅ 직접 전송 완료: 성공 {ok_n} / 시설 soft {soft_n} / 실패 {ng} "
        f"/ 전체 {len(results)}"
    )
    return results


def resend_pending_batch(
    targets: list[dict],
    report_excel_map: dict[str, str] | None = None,
    report_excel_paths: list[str] | None = None,
    *,
    mark_done: bool = True,
    progress_cb=None,
) -> list[dict]:
    """
    재전송 대상 목록을 성적서 엑셀로 다시 전송 (데이터+PDF, eco_input과 동일).
    성공 시 해당 전송로그의 재전송상태=완료 로 갱신 (다음 조회에서 제외).

    report_excel_map: {시료번호: 성적서경로}
    report_excel_paths: 파일 목록 (시료번호는 파일명에서 추출) — map과 합침
    progress_cb: 선택. (cur, total, sample_no, status) 호출.
    """
    path_map = dict(report_excel_map or {})
    if report_excel_paths:
        path_map.update(map_report_excels_by_sample(report_excel_paths))

    if not targets:
        print("▶ 재전송 대상 없음")
        return []

    total = len(targets)
    print(f"▶ 그룹웨어 기간 재전송 시작 ({total}건) — 데이터+PDF (eco_input과 동일)")
    gw_log = GroupwareRunLog()
    results = []
    for i, t in enumerate(targets, start=1):
        sno = t.get("sample_no") or ""
        src = (path_map.get(sno) or t.get("source_excel") or "").strip()
        log_path = t.get("log_path") or ""
        if progress_cb:
            try:
                progress_cb(i, total, sno, "재전송 중")
            except Exception:
                pass
        if not src or not os.path.isfile(src):
            msg = "성적서 엑셀 없음 — 파일을 추가하세요"
            print(f"  ❌ {sno}: {msg}")
            results.append({
                "sample_no": sno, "ok": False, "soft_only": False,
                "error": msg, "warnings": [], "verify_key": "",
                "log_path": log_path,
            })
            continue

        print(f"  → {sno} 재전송... excel={os.path.basename(src)}")
        row = _sync_excel_with_pdf(sno, src, gw_log)
        row["log_path"] = log_path
        ok = bool(row.get("ok"))
        soft_ok_data = bool(row.get("soft_only"))

        if ok:
            print(
                f"  ✅ {sno}: 재전송 성공 verify_key={row.get('verify_key','')}"
                f" pdf={'OK' if row.get('pdf_ok') else '없음'}"
            )
            if mark_done and log_path:
                n = mark_resend_status_in_log(
                    log_path, [sno], RESEND_STATUS_DONE,
                    note_append=f"재전송완료 {datetime.now().strftime('%Y-%m-%d %H:%M')}",
                    pdf_status="OK" if row.get("pdf_ok") else None,
                    data_status="OK" if row.get("data_ok") else None,
                )
                print(f"     ↳ 로그 갱신({n}행): {os.path.basename(log_path)} → {RESEND_STATUS_DONE}")
        elif soft_ok_data:
            print(f"  ⚠ {sno}: 데이터 OK지만 시설 soft 유지 {row.get('warnings')} — 로그는 대기로 유지")
            if mark_done and log_path:
                mark_resend_status_in_log(
                    log_path, [sno], RESEND_STATUS_PENDING,
                    note_append="; ".join(str(w) for w in (row.get("warnings") or [])),
                )
        else:
            print(f"  ❌ {sno}: {row.get('error','')} {row.get('warnings')}")

        results.append(row)
        time.sleep(1.5)

    ok_n = sum(1 for r in results if r.get("ok"))
    print(f"✅ 기간 재전송 완료: 성공 {ok_n} / 전체 {len(results)}")
    return results


def resend_from_summary_excel(
    excel_path: str,
    *,
    only_failed: bool = True,
    include_facility_soft: bool = False,
    sample_nos: list[str] | None = None,
    report_excel_map: dict[str, str] | None = None,
    mark_done: bool = True,
) -> list[dict]:
    """단일 전송엑셀 기준 재전송 (하위 호환). 성공 시 로그 재전송상태 갱신."""
    if not excel_path or not os.path.isfile(excel_path):
        raise FileNotFoundError(f"전송 엑셀 없음: {excel_path}")

    samples = _iter_summary_excel_samples(excel_path)
    want = {s.strip() for s in (sample_nos or []) if s and str(s).strip()}
    targets = []
    for s in samples:
        sno = s["sample_no"]
        if want and sno not in want:
            continue
        reason = _classify_sample_resend_need(
            s,
            include_facility_soft=include_facility_soft or bool(want),
            include_pdf_missing=True,
        )
        if want:
            s = dict(s)
            s["reason"] = reason or "지정"
            s["log_path"] = excel_path
            targets.append(s)
        elif reason:
            s = dict(s)
            s["reason"] = reason
            s["log_path"] = excel_path
            targets.append(s)
        elif not only_failed and not include_facility_soft:
            s = dict(s)
            s["reason"] = "전체"
            s["log_path"] = excel_path
            targets.append(s)

    return resend_pending_batch(
        targets,
        report_excel_map=report_excel_map,
        mark_done=mark_done,
    )


if __name__ == "__main__":
    # 기간 조회: python groupware_client.py --from 2026-08-01 --to 2026-08-05
    # 재전송:   python groupware_client.py --from ... --to ... --resend --reports "path1" "path2"
    import sys as _sys
    import argparse

    ap = argparse.ArgumentParser(description="그룹웨어 전송로그 기간 조회/재전송")
    ap.add_argument("excel", nargs="?", help="단일 그룹웨어전송_*.xlsx (구 방식)")
    ap.add_argument("--from", dest="date_from", help="시작일 YYYY-MM-DD")
    ap.add_argument("--to", dest="date_to", help="종료일 YYYY-MM-DD")
    ap.add_argument("--facility-soft", action="store_true", help="시설 soft 포함")
    ap.add_argument("--no-pdf-missing", action="store_true", help="PDF 없음(N/A) 제외")
    ap.add_argument("--resend", action="store_true", help="조회 후 재전송")
    ap.add_argument("--reports", nargs="*", default=[], help="성적서 엑셀 경로들")
    ap.add_argument("samples", nargs="*", help="시료번호(선택)")
    args = ap.parse_args()

    if args.date_from or args.date_to:
        pending = collect_pending_resends(
            args.date_from, args.date_to,
            include_facility_soft=args.facility_soft or True,
            include_pdf_missing=not args.no_pdf_missing,
        )
        for p in pending:
            print(
                f"  · {p['sample_no']}  [{p.get('reason')}]  "
                f"facility={p.get('facility_name')!r}  log={os.path.basename(p.get('log_path',''))}"
            )
        if args.resend:
            resend_pending_batch(pending, report_excel_paths=args.reports)
        raise SystemExit(0)

    if not args.excel:
        print(
            "사용법:\n"
            "  python groupware_client.py --from 2026-08-01 --to 2026-08-05\n"
            "  python groupware_client.py --from 2026-08-01 --to 2026-08-05 --resend --reports 성적서1.xlsm ...\n"
            "  python groupware_client.py <그룹웨어전송_*.xlsx> [--facility-soft] [시료번호...]"
        )
        raise SystemExit(1)

    resend_from_summary_excel(
        args.excel,
        only_failed=not args.facility_soft and not args.samples,
        include_facility_soft=args.facility_soft,
        sample_nos=args.samples or None,
        report_excel_map=map_report_excels_by_sample(args.reports) if args.reports else None,
    )
